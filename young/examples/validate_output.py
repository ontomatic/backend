#!/usr/bin/env python3
"""제품군/10종/4형식 구조와 XLSX 스키마를 검사한다.

사용법: python validate_output.py /path/to/결과폴더
시각 품질, 문서 간 값 일치, 사양 근거의 진위는 별도로 확인해야 한다.
"""
import argparse
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


TYPES = [
    "1_검사성적서", "2_부적합보고서_NCR_CAPA", "3_품질관리계획서_ControlPlan",
    "4_품질검사_체크리스트", "5_SPC관리도_XbarR", "6_협력사_품질평가서",
    "7_고객클레임_처리보고서", "8_품질회의록", "9_공정검사일지_작업일보",
    "10_수입검사성적서_IQC",
]
EXTENSIONS = ("hwpx", "docx", "xlsx", "pdf")


def schemas():
    contract = (Path(__file__).resolve().parents[1] / "references" / "qc_contract.md").read_text()
    found = {}
    for number, name, cols in re.findall(r"### 2-(\d+)\. ([^\n]+)\n```\n(.*?)\n```", contract, re.S):
        if int(number) > 10:
            continue
        title = name.split(" (")[0]
        found[title] = [x.strip() for x in cols.replace("\n", " ").split(",")]
    return found


def validate(root):
    errors = []
    schema = schemas()
    if set(schema) != set(TYPES):
        errors.append("참고 스키마에서 10종 헤더를 읽지 못했습니다")
    products = sorted(p for p in root.iterdir() if p.is_dir()) if root.is_dir() else []
    if not products:
        return [f"제품군 폴더가 없습니다: {root}"]
    for product in products:
        dirs = {p.name for p in product.iterdir() if p.is_dir()}
        if dirs != set(TYPES):
            errors.append(f"{product.name}: 문서 폴더 불일치, 누락={sorted(set(TYPES)-dirs)}, 추가={sorted(dirs-set(TYPES))}")
        for doc_type in TYPES:
            folder = product / doc_type
            if not folder.is_dir():
                continue
            stem = doc_type.split("_", 1)[1] + "_원본"
            expected = {f"{stem}.{ext}" for ext in EXTENSIONS}
            actual = {p.name for p in folder.iterdir() if p.is_file()}
            if actual != expected:
                errors.append(f"{product.name}/{doc_type}: 파일 불일치, 누락={sorted(expected-actual)}, 추가={sorted(actual-expected)}")
            for ext in EXTENSIONS:
                path = folder / f"{stem}.{ext}"
                if not path.is_file() or path.stat().st_size == 0:
                    continue
                try:
                    if ext == "xlsx":
                        wb = load_workbook(path, read_only=True, data_only=False)
                        ws = wb["데이터"] if "데이터" in wb else wb.active
                        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                        if header != schema.get(doc_type):
                            errors.append(f"{path}: XLSX 헤더/순서 불일치")
                        if ws.max_row < 2:
                            errors.append(f"{path}: 데이터 행 없음")
                        wb.close()
                    elif ext in ("hwpx", "docx"):
                        if not zipfile.is_zipfile(path) or zipfile.ZipFile(path).testzip():
                            errors.append(f"{path}: 유효한 ZIP 기반 문서가 아님")
                    elif ext == "pdf":
                        if path.open("rb").read(5) != b"%PDF-":
                            errors.append(f"{path}: PDF 서명 없음")
                except Exception as exc:
                    errors.append(f"{path}: 열기 실패 ({exc})")
    print(f"제품군 {len(products)}개, 기대 핵심 파일 {len(products)*40}개")
    return errors


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    problems = validate(args.root)
    for problem in problems:
        print("오류:", problem)
    print("구조 검사:", "통과" if not problems else f"실패 {len(problems)}건")
    sys.exit(bool(problems))
