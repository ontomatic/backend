from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape,A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak

ROOT=Path(__file__).parent
OUT=ROOT/'output'
PACKS=json.loads((ROOT/'data.json').read_text())
SCHEMAS=json.loads((ROOT/'schema.json').read_text())
FONT='QCKorean'
pdfmetrics.registerFont(TTFont(FONT,'/System/Library/Fonts/Supplemental/Arial Unicode.ttf'))

def title(name):return name.split('_',1)[1].replace('_',' ')

def overview_fields(name):
    if name.startswith('1_'):return ['행ID','검사항목','기준치','측정값','결과','비고']
    if name.startswith('2_'):return ['행ID','발견공정','부적합내용','처리판정','종결일']
    if name.startswith('3_'):return ['행ID','공정명','관리특성(사양)','규격·공차','측정방법']
    if name.startswith('4_'):return ['행ID','검사항목','판정기준','확인결과','처리방법']
    if name.startswith('5_'):return ['행ID','부분군번호','평균(X̄)','CL','UCL','LCL','이상원인·조치']
    if name.startswith('6_'):return ['행ID','협력사명','평점(100점)','등급','평가자']
    if name.startswith('7_'):return ['행ID','고객사','사고개요','원인분석','처리결과']
    if name.startswith('8_'):return ['행ID','안건','시간','기록자','결재']
    if name.startswith('9_'):return ['행ID','공정명','품질 체크포인트','확인결과','수정조치내용']
    return ['행ID','부품명','검사항목','기준치','결과','합부판정']

def full_record(row):
    return '  |  '.join(f'{k}: {v if v!="" else "[공란]"}' for k,v in row.items())

def summary_value(row,key):
    val=row[key]
    return str(val) if val!='' else '[공란]'

def source_label(row):
    note=str(row.get('비고',''))
    m=re.search(r'Specifications (\d+)쪽',note)
    return f'p.{m.group(1)}' if m else '관행/합성'

def lot_groups(rows):
    groups=[]
    for row in rows:
        lot=row['LOT No.']
        if not groups or groups[-1][0]!=lot:groups.append((lot,[]))
        groups[-1][1].append(row)
    return groups

def make_xlsx(path,pack,name,rows):
    wb=Workbook();ws=wb.active;ws.title='데이터';headers=SCHEMAS[name]
    ws.append(headers)
    for row in rows:ws.append([row[h] for h in headers])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    ws.row_dimensions[1].height=35
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='33383B');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=min(60,max(16,len(h)*2+4))
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        if row[0].row%2==0:
            for c in row:c.fill=PatternFill('solid',fgColor='F2F4F4')
    note=wb.create_sheet('안내');note.column_dimensions['A'].width=105
    for line in [f"모델: {pack['model']}",f"문서: {title(name)}",'origin=synthetic: 검사기록·운영정보는 모두 가상',f"사양 근거: {pack['source_url']}",'의도적 결함: 모델 폴더의 11_결함_주입표.xlsx 참조','SPC 관리한계: 합성 특성치의 고정 관리규칙 CL=4.00, UCL=4.15, LCL=3.85']:
        note.append([line])
    wb.save(path)

def make_docx(path,pack,name,rows):
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(29.7);sec.page_height=Cm(21);sec.top_margin=Cm(1.4);sec.bottom_margin=Cm(1.3);sec.left_margin=Cm(1.4);sec.right_margin=Cm(1.4)
    for sty in ['Normal','Title','Heading 1','Heading 2']:
        s=doc.styles[sty];s.font.name='Noto Sans Gothic';s._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans Gothic')
    doc.styles['Normal'].font.size=Pt(8)
    if name.startswith('1_'):
        groups=lot_groups(rows)
        for gi,(lot,group) in enumerate(groups):
            if gi:doc.add_page_break()
            p=doc.add_paragraph(style='Title');p.add_run(f"{pack['model']} 검사 특성 검증 기록");p.style.font.color.rgb=RGBColor(0,0,0)
            first=group[0]
            doc.add_paragraph(f"성적서 {first['성적서번호']}  |  LOT {lot}  |  검사일 {first['검사일자']}  |  표본 {first['검사수량']}/{first['총수량']} EA")
            cols=['행ID','분류','검사항목','기준치','측정값','결과','사양 근거']
            tab=doc.add_table(rows=1,cols=len(cols));tab.style='Table Grid'
            for j,h in enumerate(cols):tab.rows[0].cells[j].text=h
            for row in group:
                vals=[row['행ID'],row['분류'],row['검사항목'],row['기준치'],row['측정값'],row['결과'],source_label(row)]
                cells=tab.add_row().cells
                for j,val in enumerate(vals):cells[j].text=str(val)
            state='보류' if any(r['결과']=='NG' for r in group) else '적합'
            doc.add_paragraph(f"부적합 참조: {' / '.join(r['행ID'] for r in group if r['결과']=='NG') or '해당 없음'}  |  종합 판정: {state}  |  처리방향: {group[0]['처리방향']}")
            doc.add_paragraph(f"검사 {first['LQC검사자']}  |  OQC 확인 {first['OQC검사자']}  |  승인 {first['승인자']}")
    else:
        p=doc.add_paragraph(style='Title');p.add_run(f"{pack['model']} {title(name)}");p.style.font.color.rgb=RGBColor(0,0,0)
        doc.add_paragraph(f"문서 식별 {pack['sample_no']}  |  행 수 {len(rows)}  |  스키마 1.0  |  origin synthetic")
        doc.add_paragraph('실제 데이터시트 사양을 참조한 합성 품질관리 기록입니다. LOT·검사자·고객·협력사 등 운영정보는 가상입니다.')
        doc.add_heading('행ID별 검토 목록',level=1)
        cols=overview_fields(name);tab=doc.add_table(rows=1,cols=len(cols));tab.style='Table Grid'
        for j,h in enumerate(cols):tab.rows[0].cells[j].text=h
        for row in rows:
            cells=tab.add_row().cells
            for j,h in enumerate(cols):cells[j].text=summary_value(row,h)
    doc.add_paragraph(f"사양 근거: {pack['source_url']}")
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.name='Noto Sans Gothic';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans Gothic')
    for table in doc.tables:
        for cell in table._cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name='Noto Sans Gothic';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans Gothic')
    doc.save(path)

def make_pdf(path,pack,name,rows):
    normal=ParagraphStyle('n',fontName=FONT,fontSize=8,leading=11,spaceAfter=0)
    head=ParagraphStyle('head',parent=normal,fontSize=16,leading=20,textColor=colors.HexColor('#272B2E'),spaceAfter=8)
    sect=ParagraphStyle('sect',parent=normal,fontSize=10,leading=14,textColor=colors.HexColor('#272B2E'),spaceBefore=12,spaceAfter=5)
    small=ParagraphStyle('small',parent=normal,fontSize=7.3,leading=10.2,spaceAfter=4)
    story=[]
    if name.startswith('1_'):
        for gi,(lot,group) in enumerate(lot_groups(rows)):
            if gi:story.append(PageBreak())
            first=group[0]
            meta=Table([[Paragraph('모델코드',normal),Paragraph('LOT No.',normal),Paragraph('검사일자',normal),Paragraph('검사수량',normal)],[Paragraph(escape(pack['model']),normal),Paragraph(escape(lot),normal),Paragraph(escape(first['검사일자']),normal),Paragraph(escape(f"{first['검사수량']}/{first['총수량']} EA"),normal)]],colWidths=[195]*4)
            meta.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EAEC')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#A3AAAE')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
            story.extend([Paragraph(escape(f"{pack['model']} 검사 특성 검증 기록"),head),Paragraph(escape(f"성적서 {first['성적서번호']}  |  스키마 1.0  |  origin synthetic"),normal),Spacer(1,8),meta,Paragraph('특성 검증 내역',sect)])
            cols=['행ID','분류','검사항목','기준치','측정값','결과','사양 근거']
            widths=[80,65,95,215,190,50,85]
            content=[[Paragraph(escape(h),normal) for h in cols]]
            for row in group:
                vals=[row['행ID'],row['분류'],row['검사항목'],row['기준치'],row['측정값'],row['결과'],source_label(row)]
                content.append([Paragraph(escape(str(v)),small) for v in vals])
            table=Table(content,colWidths=widths,repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EAEC')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#A3AAAE')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F7F8F8')]),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            state='보류' if any(r['결과']=='NG' for r in group) else '적합'
            story.extend([table,Spacer(1,8),Paragraph(escape(f"부적합 참조: {' / '.join(r['행ID'] for r in group if r['결과']=='NG') or '해당 없음'}  |  종합 판정: {state}  |  처리방향: {group[0]['처리방향']}"),normal),Paragraph(escape(f"검사 {first['LQC검사자']}  |  OQC 확인 {first['OQC검사자']}  |  승인 {first['승인자']}"),normal)])
    else:
        story=[Paragraph(escape(f"{pack['model']} {title(name)}"),head),Paragraph(escape(f"문서 식별 {pack['sample_no']}  |  행 수 {len(rows)}  |  스키마 1.0  |  origin synthetic"),normal),Spacer(1,5),Paragraph('실제 데이터시트 사양을 참조한 합성 품질관리 기록입니다. LOT·검사자·고객·협력사 등 운영정보는 가상입니다.',normal),Paragraph('행ID별 검토 목록',sect)]
        cols=overview_fields(name);width=780/len(cols)
        content=[[Paragraph(escape(h),normal) for h in cols]]
        for row in rows:content.append([Paragraph(escape(summary_value(row,h)),small) for h in cols])
        table=Table(content,colWidths=[width]*len(cols),repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EAEC')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#A3AAAE')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F7F8F8')]),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
        story.append(table)
    story.append(Paragraph(escape(f"사양 근거: {pack['source_url']}"),small))
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),rightMargin=30,leftMargin=30,topMargin=30,bottomMargin=30)
    doc.build(story)

def make_hwpx(path,pack,name,rows):
    paragraphs=[f"{pack['model']} {title(name)}",f"문서 식별 {pack['sample_no']} | 행 수 {len(rows)} | 스키마 1.0 | origin synthetic",'실제 데이터시트 사양을 참조한 합성 품질관리 기록. 운영정보는 가상.','행ID별 검토 목록']
    cols=overview_fields(name)
    paragraphs.append(' | '.join(cols))
    for row in rows:paragraphs.append(' | '.join(summary_value(row,h) for h in cols))
    paragraphs.append(f"사양 근거: {pack['source_url']}")
    body=''.join(f'<hp:p id="{i}" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="{1 if i in (0,3) or p=='전체 필드 기록' else 0}"><hp:t>{escape(p)}</hp:t></hp:run></hp:p>' for i,p in enumerate(paragraphs))
    parts={
      'version.xml':'<?xml version="1.0" encoding="UTF-8"?><ha:HCFVersion xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app" targetApplication="WORDPROC" major="5" minor="1" micro="0" buildNumber="0" os="Windows"/>',
      'META-INF/container.xml':'<?xml version="1.0" encoding="UTF-8"?><ocf:container xmlns:ocf="urn:oasis:names:tc:opendocument:xmlns:container"><ocf:rootfiles><ocf:rootfile full-path="Contents/content.hpf" media-type="application/hwpml-package+xml"/></ocf:rootfiles></ocf:container>',
      'Contents/content.hpf':f'<?xml version="1.0" encoding="UTF-8"?><opf:package xmlns:opf="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/" version="1.0"><opf:metadata><dc:title>{escape(pack["model"]+" "+title(name))}</dc:title><dc:creator>온토매틱 synthetic QC</dc:creator><dc:language>ko-KR</dc:language></opf:metadata><opf:manifest><opf:item id="header" href="header.xml" media-type="application/xml"/><opf:item id="section0" href="section0.xml" media-type="application/xml"/><opf:item id="settings" href="settings.xml" media-type="application/xml"/></opf:manifest><opf:spine><opf:itemref idref="section0"/></opf:spine></opf:package>',
      'Contents/settings.xml':'<?xml version="1.0" encoding="UTF-8"?><ha:HWPApplicationSetting xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app"/>',
      'Contents/header.xml':'<?xml version="1.0" encoding="UTF-8"?><hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"><hh:beginNum page="1" footnote="1" endnote="1" pic="1" tbl="1" equation="1"/><hh:refList><hh:fontfaces><hh:fontface lang="HANGUL"><hh:font name="함초롬바탕" type="TTF"/></hh:fontface></hh:fontfaces><hh:borderFills><hh:borderFill id="0"/></hh:borderFills><hh:charProperties><hh:charPr id="0" height="850" textColor="#000000"/><hh:charPr id="1" height="1150" textColor="#272B2E"><hh:bold/></hh:charPr></hh:charProperties><hh:paraProperties><hh:paraPr id="0" align="LEFT"/></hh:paraProperties><hh:styles><hh:style id="0" type="PARA" name="바탕글" paraPrIDRef="0" charPrIDRef="0"/></hh:styles><hh:bullets/><hh:numberings/></hh:refList></hh:head>',
      'Contents/section0.xml':f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">{body}</hs:sec>',
      'Preview/PrvText.txt':'\n'.join(paragraphs)
    }
    with zipfile.ZipFile(path,'w') as z:
        z.writestr('mimetype','application/hwp+zip',compress_type=zipfile.ZIP_STORED)
        for key,value in parts.items():z.writestr(key,value,compress_type=zipfile.ZIP_DEFLATED)

def defects_xlsx(path,pack):
    wb=Workbook();ws=wb.active;ws.title='결함_주입표';ws.append(['시트','행ID','결함유형','설명'])
    for d in pack['defects']:ws.append(list(d.values()))
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='33383B');c.font=Font(color='FFFFFF',bold=True)
    for col,width in zip('ABCD',[38,23,28,80]):ws.column_dimensions[col].width=width
    wb.save(path)

def readme(path,pack):
    lines=[f"# {pack['model']} QC 합성데이터",'',f"원본 데이터시트: {pack['source_url']}",'','10종 문서 × 4형식(.hwpx/.docx/.xlsx/.pdf) = 40개 파일. 11_결함_주입표.xlsx는 의도적 결함의 정답표입니다.','',"1_검사성적서의 기준치 중 외관을 제외한 값은 원본 데이터시트 Specifications에서 확인했습니다. 각 레코드 비고에 페이지를 기록했습니다. 운영정보와 2~10종 문서값은 가상이며 실제 한화비전 QC·인사·협력사·고객 데이터가 아닙니다.",'',"SPC 관리한계는 합성 특성치 '조립 정렬 간격(mm)'에 대해 CL=4.00, UCL=4.15, LCL=3.85로 고정했습니다. 평균과 범위는 5개 측정값에서 계산했습니다.",'',"HWPX는 공개 XML 패키지 형식으로 작성했으며 한컴오피스 직접 열기 검증은 수행하지 못했습니다.",'',"문서 유형은 10종이며 IQC가 10번째입니다. 이전 9종에서 10종으로 변경된 상세 이력은 첨부 스킬에 없어 추정하지 않았습니다."]
    if pack['omissions']:lines+=['','## 제외한 사양']+[f'- {x}' for x in pack['omissions']]
    path.write_text('\n'.join(lines),encoding='utf-8')

def main():
    count=0
    for pack in PACKS:
        sample=OUT/f"{pack['sample_no']}_{pack['model']}";sample.mkdir(parents=True,exist_ok=True)
        for name,rows in pack['data'].items():
            folder=sample/name;folder.mkdir(exist_ok=True)
            stem=folder/f"{pack['model']}_{name}"
            make_xlsx(stem.with_suffix('.xlsx'),pack,name,rows)
            make_docx(stem.with_suffix('.docx'),pack,name,rows)
            make_pdf(stem.with_suffix('.pdf'),pack,name,rows)
            make_hwpx(stem.with_suffix('.hwpx'),pack,name,rows)
            count+=4
        defects_xlsx(sample/'11_결함_주입표.xlsx',pack)
        readme(sample/'읽어주세요.md',pack)
        print(pack['sample_no'],pack['model'],'done',flush=True)
    print('files',count)

if __name__=='__main__':main()
