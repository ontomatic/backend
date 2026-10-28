"""독립 미사용 데이터시트 두 건에 대한 스킬 실행 검증."""
import json
import io
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "runtime"))
from generate_data import make_model, validate  # noqa: E402
import render_files  # noqa: E402
from render_files import make_xlsx, make_docx, make_pdf, make_hwpx  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402
from reportlab.pdfbase import pdfmetrics  # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont  # noqa: E402


def source_page(row):
    note = str(row.get("비고", ""))
    found = re.search(r"(\d+)쪽", note)
    return f"p.{found.group(1)}" if found else "합성 기준"


render_files.source_label = source_page


def item(name, value, page, section="Specifications"):
    return {"name": name, "value": value, "page": page, "section": section}


def use_korean_font(path):
    doc = Document(path)
    doc.add_paragraph("합성 데이터: 실제 검사·거래·인사 기록이 아닙니다.")
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "AppleGothic"
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "AppleGothic")
    for paragraph in list(doc.paragraphs) + [p for t in doc.tables for c in t._cells for p in c.paragraphs]:
        for run in paragraph.runs:
            run.font.name = "AppleGothic"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "AppleGothic")
    doc.save(path)


def add_pdf_notice(path):
    reader = PdfReader(str(path))
    writer = PdfWriter()
    font_name = "QCTestKorean"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"))
    for page in reader.pages:
        packet = io.BytesIO()
        c = canvas.Canvas(packet, pagesize=(float(page.mediabox.width), float(page.mediabox.height)))
        c.setFont(font_name, 8)
        c.drawString(34, 15, "합성 데이터: 실제 검사·거래·인사 기록이 아닙니다.")
        c.save()
        packet.seek(0)
        page.merge_page(PdfReader(packet).pages[0])
        writer.add_page(page)
    with path.open("wb") as stream:
        writer.write(stream)


SOURCES = [
    {
        "sample_no": "18", "category": "18-네트워크-불렛-XNO-8080R", "model": "XNO-8080R",
        "source_id": "13wWwZF8t7Wmc3bjVEawBZkKTV0_PV-9A",
        "source_url": "https://drive.google.com/file/d/13wWwZF8t7Wmc3bjVEawBZkKTV0_PV-9A/view",
        "omissions": ["없음: 아래 QC 검사 특성은 모두 원본 Specifications에서 확인"],
        "conflicts": ["최저조도: 1쪽 Key Features '0Lux (IR LED On)'는 IR 조사 조건이고, 2쪽 Specifications 컬러 0.07Lux/흑백 0.007Lux(1/30초, 30IRE)를 일반 최저조도 기준으로 채택"],
        "items": [
            item("외관", "흠집·변형 없음", None, "합성 검사 기준·원본 사양 아님"),
            item("해상도", "2560x1920 (5MP)", 2),
            item("화각수평", "100.2°~38.7°", 2),
            item("줌배율", "광학 2.5x", 2),
            item("최저조도", "컬러 0.07Lux / 흑백 0.007Lux (1/30초, 30IRE)", 2),
            item("초점거리", "3.7~9.4mm", 2),
            item("방수내충격등급", "IP66 / IP67 / IK10", 4, "Certifications & Standards"),
            item("전원", "PoE(IEEE802.3af, Class3) / 12VDC / 24VAC", 4),
            item("동작온도", "-40°C~+55°C", 4, "동작 조건"),
            item("습도", "0~90% RH(non-condensing)", 4, "동작 조건"),
            item("소비전력", "PoE Max 12.95W / 12VDC Max 12.5W / 24VAC Max 14.5W", 4),
            item("색상", "다크그레이", 4),
            item("재질", "알루미늄", 4),
            item("제품사이즈", "Ø91x368.6mm", 4),
            item("제품무게", "2180g", 4),
        ],
    },
    {
        "sample_no": "19", "category": "19-아날로그-돔-HCD-6080R", "model": "HCD-6080R",
        "source_id": "1EHSqKOI2A5lai0xAtrinY0kdCSF6NC94",
        "source_url": "https://drive.google.com/file/d/1EHSqKOI2A5lai0xAtrinY0kdCSF6NC94/view",
        "omissions": ["방수내충격등급: 해당 데이터시트에서 IP/IK 인증 근거 미확인, 항목 제외"],
        "conflicts": ["최저조도: 1쪽 Key Features '0Lux (IR LED On)'는 IR 조사 조건이고, 2쪽 Specifications 컬러 0.11Lux/흑백 0.011Lux(1/30초, 30IRE)를 일반 최저조도 기준으로 채택"],
        "items": [
            item("외관", "흠집·변형 없음", None, "합성 검사 기준·원본 사양 아님"),
            item("해상도", "1920x1080 (2MP)", 2),
            item("화각수평", "110.2°~32°", 2),
            item("줌배율", "광학 3.1x", 2),
            item("최저조도", "컬러 0.11Lux / 흑백 0.011Lux (1/30초, 30IRE)", 2),
            item("초점거리", "3.2~10mm", 2),
            item("전원", "24VAC / 12VDC", 3),
            item("동작온도", "-10°C~+55°C", 3, "동작 조건"),
            item("습도", "0~90% RH(non-condensing)", 3, "동작 조건"),
            item("소비전력", "12VDC 6W / 24VAC 6W", 3),
            item("색상", "아이보리", 3),
            item("재질", "플라스틱(Polycarbonate)", 3),
            item("제품사이즈", "ø119.8x98.5mm", 3),
            item("제품무게", "330g", 3),
        ],
    },
]


def main():
    out = ROOT / "result_D"
    count = 0
    reports = []
    for source in SOURCES:
        pack = make_model(source)
        validate(pack)
        # 요약 문구와 조건부 사양표의 차이는 실제 전달 문서의 비고에도 남긴다.
        for row in pack["data"]["1_검사성적서"]:
            if row["검사항목"] == "최저조도":
                row["비고"] += "; " + source["conflicts"][0]
        product = out / source["category"]
        for doc_type, rows in pack["data"].items():
            folder = product / doc_type
            folder.mkdir(parents=True, exist_ok=True)
            stem = folder / (doc_type.split("_", 1)[1] + "_원본")
            make_xlsx(stem.with_suffix(".xlsx"), pack, doc_type, rows)
            make_docx(stem.with_suffix(".docx"), pack, doc_type, rows)
            use_korean_font(stem.with_suffix(".docx"))
            make_pdf(stem.with_suffix(".pdf"), pack, doc_type, rows)
            add_pdf_notice(stem.with_suffix(".pdf"))
            make_hwpx(stem.with_suffix(".hwpx"), pack, doc_type, rows)
            count += 4
        reports.append({"category": source["category"], "model": source["model"], "source_url": source["source_url"], "items": source["items"], "conflicts": source["conflicts"], "omissions": source["omissions"], "rows": {k: len(v) for k, v in pack["data"].items()}, "defects": pack["defects"]})
    (ROOT / "validation_evidence.json").write_text(json.dumps(reports, ensure_ascii=False, indent=2))
    print(f"products={len(SOURCES)} files={count}")


if __name__ == "__main__":
    main()
