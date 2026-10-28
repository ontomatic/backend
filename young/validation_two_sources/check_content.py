"""두 모델 시험 결과의 값·근거·형식 간 핵심 내용 검사."""
import json
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from build_validation import ROOT, SOURCES


def main():
    errors = []
    count = 0
    for source in SOURCES:
        product = ROOT / "result_D" / source["category"]
        for folder in product.iterdir():
            if not folder.is_dir():
                continue
            stem = folder / (folder.name.split("_", 1)[1] + "_원본")
            wb = load_workbook(stem.with_suffix(".xlsx"), read_only=True, data_only=True)
            ws = wb["데이터"]
            headers = [c.value for c in ws[1]]
            rows = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
            ids = {str(r["행ID"]) for r in rows}
            pdf = " ".join(page.extract_text() or "" for page in PdfReader(stem.with_suffix(".pdf")).pages)
            doc = Document(stem.with_suffix(".docx"))
            docx = " ".join(p.text for p in doc.paragraphs) + " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            with zipfile.ZipFile(stem.with_suffix(".hwpx")) as archive:
                hwpx = archive.read("Preview/PrvText.txt").decode()
            for label, text in [("pdf", pdf), ("docx", docx), ("hwpx", hwpx)]:
                missing = ids - {rid for rid in ids if rid in text}
                if missing:
                    errors.append(f"{source['model']}/{folder.name}/{label}: 행ID {len(missing)}개 누락")
            if "합성 데이터" not in pdf or "합성 데이터" not in docx:
                errors.append(f"{source['model']}/{folder.name}: PDF 또는 DOCX 합성 안내 누락")
            if folder.name.startswith("1_"):
                items = {r["검사항목"].replace(" ", ""): r for r in rows if r["LOT No."].endswith("001")}
                for item in source["items"]:
                    row = items.get(item["name"])
                    if row is None:
                        errors.append(f"{source['model']}: {item['name']} 검사항목 누락")
                    elif row["기준치"] != item["value"]:
                        errors.append(f"{source['model']}: {item['name']} 기준치 불일치")
                    if item["page"] and f"p.{item['page']}" not in pdf:
                        errors.append(f"{source['model']}: {item['name']} PDF 근거 페이지 누락")
                if source["model"] == "HCD-6080R" and "방수내충격등급" in items:
                    errors.append("HCD-6080R: 원본에 없는 IP/IK 사양이 생성됨")
                if source["conflicts"][0] not in " ".join(str(r.get("비고")) for r in rows):
                    errors.append(f"{source['model']}: 최저조도 상충 설명 누락")
            wb.close()
            count += 1
    report = {"documents": count, "files": count * 4, "machine_content_errors": errors, "manual_visual_result": "DOCX 렌더링 시 한글 누락: 미통과", "hwpx_application_result": "한컴오피스 직접 열기 미검증"}
    (ROOT / "validation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(bool(errors))


if __name__ == "__main__":
    main()
