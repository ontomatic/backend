from io import BytesIO

from backend.parsers.documents.interface import DocumentBlock, ParsedDocument, SourceLocation, TableBlock, TextBlock


def parse_xlsx(source: BytesIO) -> ParsedDocument:
  from openpyxl import load_workbook
  from openpyxl.utils.cell import get_column_letter

  wb = load_workbook(source, read_only=True, data_only=True)
  blocks: list[DocumentBlock] = []
  try:
    for sheet in wb.worksheets:
      sheet.reset_dimensions()
      min_row = min_col = None
      max_row = max_col = 0

      for row in sheet.iter_rows():
        for cell in row:
          if cell.value is None:
            continue

          min_row = cell.row if min_row is None else min(min_row, cell.row)
          min_col = cell.column if min_col is None else min(min_col, cell.column)
          max_row = max(max_row, cell.row)
          max_col = max(max_col, cell.column)

      if min_row is None or min_col is None:
        continue

      cell_range = f'{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}'
      iter_rows = sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)
      rows = [[str(value) if value is not None else None for value in row] for row in iter_rows]
      blocks.append(TableBlock(rows=rows, location=SourceLocation(sheet_name=sheet.title, cell_range=cell_range)))
  finally:
    wb.close()

  return ParsedDocument(format='xlsx', blocks=blocks)
