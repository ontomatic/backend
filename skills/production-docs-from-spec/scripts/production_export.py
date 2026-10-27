"""생산관리 원본 xlsx 를 PDF·docx·hwpx 로도 내보낸다 (spec_to_production.py 가 부른다).

세 형식 모두 이미 저장한 <문서>_원본.xlsx 를 다시 읽어 만든다. 그래서 값은 xlsx 와 같고,
check_formats() 가 형식마다 글자를 다시 뽑아 xlsx 와 칸 단위로 맞는지 검사한다.
- PDF : LibreOffice 로 xlsx 를 그대로 인쇄 (가로 A4, 한 쪽 너비 맞춤, 표 머리글 매 쪽 반복)
- docx: python-docx 로 같은 머리(제목·회사·출력일·조회조건) + 표
- hwpx: OWPML(zip + XML) 을 직접 쓴다. 한글 없이 만든 파일이라 한글에서 열어 확인해야 한다
"""
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import get_column_letter

SOFFICE_PLACES = [r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                  "/Applications/LibreOffice.app/Contents/MacOS/soffice", "/usr/bin/soffice", "/usr/local/bin/soffice"]
FONT = "맑은 고딕"
PAGE_W_MM, PAGE_H_MM, MARGIN_MM = 297, 210, 10          # 가로 A4
USABLE_MM = PAGE_W_MM - 2 * MARGIN_MM
CELL_PAD_MM = 1.0                                         # docx·hwpx 칸 안쪽 좌우 여백


def find_soffice():
    """LibreOffice 실행 파일: SOFFICE 환경변수 → PATH → 흔한 설치 위치 순서로 찾는다. 없으면 None."""
    for c in [os.environ.get("SOFFICE"), shutil.which("soffice"), shutil.which("libreoffice")] + SOFFICE_PLACES:
        if c and Path(c).is_file():
            return c
    return None


# ═════════════════════════ 1. xlsx 한 장 → 공통 모양 ═════════════════════════
def _text(c):
    """엑셀이 화면에 보여 주는 글자 그대로 (General 숫자는 군더더기 0 없이)."""
    v = c.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        if c.number_format == "0.0":
            return f"{v:.1f}"
        return f"{v:g}" if isinstance(v, float) else str(v)
    return str(v)


def _fill(c):
    if c.fill is None or c.fill.fill_type != "solid":
        return None
    rgb = c.fill.fgColor.rgb
    return rgb[-6:].upper() if isinstance(rgb, str) else None


def _is_header(ws, r):
    cells = [ws.cell(r, c) for c in range(1, ws.max_column + 1)]
    return all(c.value is not None and c.font is not None and c.font.bold and _fill(c) for c in cells)


def read_sheet(ws):
    """제목·회사·출력일·머리 정보 줄·표 머리글·데이터 줄을 뽑는다. 표 머리글 = 칸이 모두 차 있고 굵고 배경색 있는 첫 줄."""
    hdr = next(r for r in range(3, ws.max_row + 1) if _is_header(ws, r))
    ncol = ws.max_column
    printed = next((_text(ws.cell(2, c)) for c in range(ncol, 1, -1) if ws.cell(2, c).value is not None), "")
    info = []                                       # 머리 정보 줄: [(글자, 항목이름인가)]
    for r in range(3, hdr):
        cells = [(ws.cell(r, c)) for c in range(1, ncol + 1) if ws.cell(r, c).value is not None]
        if cells:
            info.append([(_text(c), bool(c.font and c.font.bold and _fill(c))) for c in cells])
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        cells = [ws.cell(r, c) for c in range(1, ncol + 1)]
        if all(c.value is None for c in cells):
            break
        rows.append([dict(text=_text(c), align=(c.alignment.horizontal if c.alignment else None) or ("right" if isinstance(c.value, (int, float)) else "left"),
                          fill=_fill(c), bold=bool(c.font and c.font.bold)) for c in cells])
    return dict(name=ws.title, title=_text(ws["A1"]), company=_text(ws["A2"]), printed=printed, info=info,
                header=[_text(ws.cell(hdr, c)) for c in range(1, ncol + 1)], header_fill=_fill(ws.cell(hdr, 1)), hdr_row=hdr,
                widths=[ws.column_dimensions[get_column_letter(c)].width or 8.43 for c in range(1, ncol + 1)], rows=rows)


def read_book(path):
    wb = openpyxl.load_workbook(path)
    try:
        return [read_sheet(ws) for ws in wb.worksheets]
    finally:
        wb.close()


# ═════════════════════════ 2. 칸 너비·글자 크기 (docx·hwpx 공통) ═════════════════════════
def _text_mm(s, pt):
    """맑은 고딕 기준 대략 글자 폭(mm): 한글 1em, 대문자 0.65em, 그 밖 0.55em."""
    em = pt * 0.3528
    return sum(em * (1.0 if ord(ch) >= 0x1100 else 0.65 if ch.isupper() else 0.55) for ch in s)


def layout(sheet):
    """긴 칸만 접히도록 칸 너비를 나눈다: 글자 크기를 9pt 부터 줄여 가며, 접히는 칸도 25mm 는 남는 크기를 고른다."""
    texts = [[sheet["header"][i]] + [r[i]["text"] for r in sheet["rows"]] for i in range(len(sheet["header"]))]
    for pt in (9, 8.5, 8, 7.5, 7, 6.5):
        nat = [max(_text_mm(t, pt) for t in col) + 2 * CELL_PAD_MM + 1 for col in texts]
        hdr = [_text_mm(h, pt) + 2 * CELL_PAD_MM + 1 for h in sheet["header"]]
        if sum(nat) <= USABLE_MM:
            k = USABLE_MM / sum(nat)
            return pt, [w * k for w in nat]
        lo, hi = 0.0, max(nat)                      # 상한 cap 을 찾는다: sum(min(nat, cap)) == USABLE
        for _ in range(40):
            cap = (lo + hi) / 2
            lo, hi = (cap, hi) if sum(min(n, max(cap, h)) for n, h in zip(nat, hdr)) < USABLE_MM else (lo, cap)
        if lo >= 25 or pt == 6.5:
            ws_ = [min(n, max(lo, h)) for n, h in zip(nat, hdr)]
            k = USABLE_MM / sum(ws_)
            return pt, [w * k for w in ws_]


# ═════════════════════════ 3. PDF (LibreOffice) ═════════════════════════
def to_pdf(xlsx_paths):
    """xlsx 여러 개를 한 번에 PDF 로 (각 xlsx 옆에 같은 이름 .pdf). 사용자가 연 LibreOffice 와 섞이지 않게 따로 된 프로필을 쓴다."""
    xlsx_paths = [Path(p) for p in xlsx_paths]
    with tempfile.TemporaryDirectory() as tmp:
        prof = Path(tmp, "profile").as_uri()
        outdir = Path(tmp, "out")
        soffice = find_soffice()
        if not soffice:
            raise RuntimeError("LibreOffice(soffice) 를 찾지 못해 PDF 를 만들 수 없습니다.")
        r = subprocess.run([soffice, f"-env:UserInstallation={prof}", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(outdir)]
                           + [str(p) for p in xlsx_paths], capture_output=True, timeout=600)
        for p in xlsx_paths:
            made = outdir / (p.stem + ".pdf")
            if not made.exists():
                raise RuntimeError(f"PDF 변환 실패: {p} ({r.stderr.decode(errors='ignore')[:300]})")
            shutil.move(str(made), p.with_suffix(".pdf"))


# ═════════════════════════ 4. docx ═════════════════════════
def to_docx(sheets, path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    AL = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}

    def shade(cell, rgb):
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), rgb)
        cell._tc.get_or_add_tcPr().append(sh)

    def put(cell, text, pt, align="left", bold=False, fill=None):
        p = cell.paragraphs[0]
        p.alignment = AL.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text)
        run.font.size = Pt(pt)
        run.bold = bold
        if fill:
            shade(cell, fill)

    def grid(tbl, color="999999"):
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), color)
            borders.append(b)
        tbl._tbl.tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("left", CELL_PAD_MM), ("right", CELL_PAD_MM), ("top", 0.3), ("bottom", 0.3)):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(round(v * 56.7)))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tbl._tbl.tblPr.append(mar)

    def widths(tbl, mm):
        tbl.autofit = False
        for col, w in zip(tbl.columns, mm):                    # 표 격자(gridCol) 너비: LibreOffice·한컴은 이것을 본다
            col.width = Mm(w)
        for row in tbl.rows:
            for c, w in zip(row.cells, mm):
                c.width = Mm(w)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(PAGE_W_MM), Mm(PAGE_H_MM)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(MARGIN_MM)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(9)
    st.paragraph_format.space_after = Pt(0)
    for n, sh in enumerate(sheets):
        p = doc.add_paragraph()
        if n:
            p.add_run().add_break(WD_BREAK.PAGE)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(sh["title"])
        r.bold, r.font.size, r.font.color.rgb = True, Pt(16), RGBColor(0x1F, 0x38, 0x64)
        p.paragraph_format.space_after = Pt(6)
        p = doc.add_paragraph()                                  # 회사 ······ 출력일 (오른쪽 끝)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(USABLE_MM), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(sh["company"])
        r.bold, r.font.size = True, Pt(9)
        r = p.add_run("\t" + sh["printed"])
        r.font.size = Pt(9)
        for line in sh["info"]:
            if any(lab for _, lab in line):                      # BOM 처럼 '항목 | 값' 이 번갈아 있는 줄은 작은 표로
                t = doc.add_table(rows=1, cols=len(line))
                grid(t)
                for c, (txt, lab) in zip(t.rows[0].cells, line):
                    put(c, txt, 9, "center" if lab else "left", lab, "EDEDED" if lab else None)
                widths(t, [max(_text_mm(x, 9) + 2 * CELL_PAD_MM + 2, 22) for x, _ in line])
            else:
                p = doc.add_paragraph()
                r = p.add_run("  ".join(t for t, _ in line))
                r.italic, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0x59, 0x59, 0x59)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        pt, mm = layout(sh)
        t = doc.add_table(rows=1 + len(sh["rows"]), cols=len(sh["header"]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        grid(t)
        for c, h in zip(t.rows[0].cells, sh["header"]):
            put(c, h, pt, "center", True, sh["header_fill"])
        trPr = t.rows[0]._tr.get_or_add_trPr()                   # 머리글 줄을 쪽마다 반복
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        trPr.append(rep)
        for row, cells in zip(t.rows[1:], sh["rows"]):
            for c, d in zip(row.cells, cells):
                put(c, d["text"], pt, d["align"], d["bold"], d["fill"])
        widths(t, mm)
    doc.save(path)


# ═════════════════════════ 5. hwpx (OWPML 직접 쓰기) ═════════════════════════
HU = 7200 / 25.4                                          # 1mm = 283.46 HWPUNIT
NS = ('xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
      'xmlns:hp10="http://www.hancom.co.kr/hwpml/2016/paragraph" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
      'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core" xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" '
      'xmlns:hhs="http://www.hancom.co.kr/hwpml/2011/history" xmlns:hm="http://www.hancom.co.kr/hwpml/2011/master-page" '
      'xmlns:hpf="http://www.hancom.co.kr/schema/2011/hpf" xmlns:dc="http://purl.org/dc/elements/1.1/" '
      'xmlns:opf="http://www.idpf.org/2007/opf/" xmlns:ooxmlchart="http://www.hancom.co.kr/hwpml/2016/ooxmlchart" '
      'xmlns:hwpunitchar="http://www.hancom.co.kr/hwpml/2016/HwpUnitChar" xmlns:epub="http://www.idpf.org/2007/ops" '
      'xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0"')
XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'
LANGS = ("hangul", "latin", "hanja", "japanese", "other", "symbol", "user")


class _Reg:
    """같은 모양이면 같은 번호를 돌려주는 글자·문단·테두리 모양 목록."""

    def __init__(self, first):
        self.items = list(first)

    def __call__(self, key):
        if key not in self.items:
            self.items.append(key)
        return self.items.index(key)


def _char_xml(i, key):
    pt, bold, italic, color = key
    each = lambda v: " ".join(f'{l}="{v}"' for l in LANGS)
    return (f'<hh:charPr id="{i}" height="{round(pt * 100)}" textColor="#{color}" shadeColor="none" useFontSpace="0" useKerning="0" symMark="NONE" borderFillIDRef="2">'
            f'<hh:fontRef {each(0)}/><hh:ratio {each(100)}/><hh:spacing {each(0)}/><hh:relSz {each(100)}/><hh:offset {each(0)}/>'
            + ("<hh:italic/>" if italic else "") + ("<hh:bold/>" if bold else "")
            + '<hh:underline type="NONE" shape="SOLID" color="#000000"/><hh:strikeout shape="NONE" color="#000000"/><hh:outline type="NONE"/>'
              '<hh:shadow type="NONE" color="#C0C0C0" offsetX="10" offsetY="10"/></hh:charPr>')


def _para_xml(i, key):
    align, spacing = key
    m = ('<hh:margin><hc:intent value="0" unit="HWPUNIT"/><hc:left value="0" unit="HWPUNIT"/><hc:right value="0" unit="HWPUNIT"/>'
         '<hc:prev value="0" unit="HWPUNIT"/><hc:next value="0" unit="HWPUNIT"/></hh:margin>'
         f'<hh:lineSpacing type="PERCENT" value="{spacing}" unit="HWPUNIT"/>')
    return (f'<hh:paraPr id="{i}" tabPrIDRef="0" condense="0" fontLineHeight="0" snapToGrid="1" suppressLineNumbers="0" checked="0" textDir="LTR">'
            f'<hh:align horizontal="{align}" vertical="BASELINE"/><hh:heading type="NONE" idRef="0" level="0"/>'
            '<hh:breakSetting breakLatinWord="KEEP_WORD" breakNonLatinWord="BREAK_WORD" widowOrphan="0" keepWithNext="0" keepLines="0" pageBreakBefore="0" lineWrap="BREAK"/>'
            '<hh:autoSpacing eAsianEng="0" eAsianNum="0"/>'
            f'<hp:switch><hp:case hp:required-namespace="http://www.hancom.co.kr/hwpml/2016/HwpUnitChar">{m}</hp:case><hp:default>{m}</hp:default></hp:switch>'
            '<hh:border borderFillIDRef="2" offsetLeft="0" offsetRight="0" offsetTop="0" offsetBottom="0" connect="0" ignoreMargin="0"/></hh:paraPr>')


def _border_xml(i, key):
    line, fill = key                                       # line: None(선 없음) / 색
    side = lambda s: (f'<hh:{s} type="SOLID" width="0.12 mm" color="#{line}"/>' if line else f'<hh:{s} type="NONE" width="0.1 mm" color="#000000"/>')
    brush = f'<hc:fillBrush><hc:winBrush faceColor="{"#" + fill if fill else "none"}" hatchColor="#999999" alpha="0"/></hc:fillBrush>' if (fill or i == 2) else ""
    return (f'<hh:borderFill id="{i}" threeD="0" shadow="0" centerLine="NONE" breakCellSeparateLine="0">'
            '<hh:slash type="NONE" Crooked="0" isCounter="0"/><hh:backSlash type="NONE" Crooked="0" isCounter="0"/>'
            + side("leftBorder") + side("rightBorder") + side("topBorder") + side("bottomBorder")
            + '<hh:diagonal type="SOLID" width="0.1 mm" color="#000000"/>' + brush + '</hh:borderFill>')


def _header_xml(chars, paras, borders):
    font = (f'<hh:font id="0" face="{FONT}" type="TTF" isEmbedded="0"><hh:typeInfo familyType="FCAT_GOTHIC" weight="6" proportion="4" contrast="0" '
            'strokeVariation="1" armStyle="1" letterform="1" midline="1" xHeight="1"/></hh:font>')
    faces = "".join(f'<hh:fontface lang="{l.upper()}" fontCnt="1">{font}</hh:fontface>' for l in LANGS)
    return (XML_HEAD + f'<hh:head {NS} version="1.5" secCnt="1"><hh:beginNum page="1" footnote="1" endnote="1" pic="1" tbl="1" equation="1"/><hh:refList>'
            f'<hh:fontfaces itemCnt="7">{faces}</hh:fontfaces>'
            f'<hh:borderFills itemCnt="{len(borders.items) - 1}">' + "".join(_border_xml(i, k) for i, k in enumerate(borders.items) if i) + '</hh:borderFills>'
            f'<hh:charProperties itemCnt="{len(chars.items)}">' + "".join(_char_xml(i, k) for i, k in enumerate(chars.items)) + '</hh:charProperties>'
            '<hh:tabProperties itemCnt="1"><hh:tabPr id="0" autoTabLeft="0" autoTabRight="0"/></hh:tabProperties>'
            f'<hh:paraProperties itemCnt="{len(paras.items)}">' + "".join(_para_xml(i, k) for i, k in enumerate(paras.items)) + '</hh:paraProperties>'
            '<hh:styles itemCnt="1"><hh:style id="0" type="PARA" name="바탕글" engName="Normal" paraPrIDRef="0" charPrIDRef="0" nextStyleIDRef="0" langID="1042" lockForm="0"/></hh:styles>'
            '</hh:refList><hh:compatibleDocument targetProgram="HWP201X"><hh:layoutCompatibility/></hh:compatibleDocument>'
            '<hh:docOption><hh:linkinfo path="" pageInherit="0" footnoteInherit="0"/></hh:docOption><hh:trackchageConfig flags="56"/></hh:head>')


def _sec_pr():
    pbf = "".join(f'<hp:pageBorderFill type="{t}" borderFillIDRef="1" textBorder="PAPER" headerInside="0" footerInside="0" fillArea="PAPER">'
                  '<hp:offset left="1417" right="1417" top="1417" bottom="1417"/></hp:pageBorderFill>' for t in ("BOTH", "EVEN", "ODD"))
    note = ('<hp:autoNumFormat type="DIGIT" userChar="" prefixChar="" suffixChar=")" supscript="0"/>'
            '<hp:noteLine length="-1" type="SOLID" width="0.12 mm" color="#000000"/><hp:noteSpacing betweenNotes="283" belowLine="567" aboveLine="850"/>'
            '<hp:numbering type="CONTINUOUS" newNum="1"/>')
    mg = round(MARGIN_MM * HU)
    return ('<hp:secPr id="" textDirection="HORIZONTAL" spaceColumns="1134" tabStop="8000" tabStopVal="4000" tabStopUnit="HWPUNIT" outlineShapeIDRef="1" '
            'memoShapeIDRef="0" textVerticalWidthHead="0" masterPageCnt="0"><hp:grid lineGrid="0" charGrid="0" wonggojiFormat="0"/>'
            '<hp:startNum pageStartsOn="BOTH" page="0" pic="0" tbl="0" equation="0"/>'
            '<hp:visibility hideFirstHeader="0" hideFirstFooter="0" hideFirstMasterPage="0" border="SHOW_ALL" fill="SHOW_ALL" hideFirstPageNum="0" hideFirstEmptyLine="0" showLineNumber="0"/>'
            '<hp:lineNumberShape restartType="0" countBy="0" distance="0" startNumber="0"/>'
            # 한글은 종이 크기를 세로 기준으로 적고 landscape="NARROWLY" 로 눕힌다 (WIDELY = 세로)
            f'<hp:pagePr landscape="NARROWLY" width="{round(PAGE_H_MM * HU)}" height="{round(PAGE_W_MM * HU)}" gutterType="LEFT_ONLY">'
            f'<hp:margin header="1417" footer="1417" gutter="0" left="{mg}" right="{mg}" top="{mg}" bottom="{mg}"/></hp:pagePr>'
            f'<hp:footNotePr>{note}<hp:placement place="EACH_COLUMN" beneathText="0"/></hp:footNotePr>'
            f'<hp:endNotePr>{note}<hp:placement place="END_OF_DOCUMENT" beneathText="0"/></hp:endNotePr>{pbf}</hp:secPr>'
            '<hp:ctrl><hp:colPr id="" type="NEWSPAPER" layout="LEFT" colCount="1" sameSz="1" sameGap="0"/></hp:ctrl>')


def to_hwpx(sheets, path):
    chars = _Reg([(10, False, False, "000000")])
    paras = _Reg([("JUSTIFY", 160)])
    borders = _Reg([None, (None, None), (None, None)])         # 0 은 비움, 1 = 쪽 테두리(선 없음), 2 = 글자 배경(없음)
    AL = {"center": "CENTER", "right": "RIGHT"}
    ids = iter(range(1, 10 ** 6))

    def para(runs, align="LEFT", spacing=130, page_break=False):
        body = "".join(f'<hp:run charPrIDRef="{chars(ck)}">{x}</hp:run>' for ck, x in runs)
        return (f'<hp:p id="{next(ids)}" paraPrIDRef="{paras((align, spacing))}" styleIDRef="0" pageBreak="{int(page_break)}" columnBreak="0" merged="0">{body}</hp:p>')

    def t(s):
        return f"<hp:t>{escape(s)}</hp:t>" if s else "<hp:t/>"

    def table(rows, widths_mm, line="999999", repeat=False, pad=CELL_PAD_MM):
        """rows: [[(글자, charKey, 정렬, 배경색)]]. 쪽에서 나뉘어 흐르는 표 (글자처럼 취급 아님)."""
        ws = [round(w * HU) for w in widths_mm]
        rh = 900
        pd = round(pad * HU)
        trs = []
        for ri, row in enumerate(rows):
            tcs = []
            for ci, (txt, ck, al, fill) in enumerate(row):
                tcs.append(f'<hp:tc name="" header="{int(repeat and ri == 0)}" hasMargin="0" protect="0" editable="0" dirty="0" borderFillIDRef="{borders((line, fill))}">'
                           '<hp:subList id="" textDirection="HORIZONTAL" lineWrap="BREAK" vertAlign="CENTER" linkListIDRef="0" linkListNextIDRef="0" '
                           'textWidth="0" textHeight="0" hasTextRef="0" hasNumRef="0">' + para([(ck, t(txt))], AL.get(al, "LEFT"), 120) + '</hp:subList>'
                           f'<hp:cellAddr colAddr="{ci}" rowAddr="{ri}"/><hp:cellSpan colSpan="1" rowSpan="1"/><hp:cellSz width="{ws[ci]}" height="{rh}"/>'
                           f'<hp:cellMargin left="{pd}" right="{pd}" top="85" bottom="85"/></hp:tc>')
            trs.append("<hp:tr>" + "".join(tcs) + "</hp:tr>")
        tbl = (f'<hp:tbl id="{next(ids)}" zOrder="0" numberingType="TABLE" textWrap="TOP_AND_BOTTOM" textFlow="BOTH_SIDES" lock="0" dropcapstyle="None" '
               f'pageBreak="CELL" repeatHeader="{int(repeat)}" rowCnt="{len(rows)}" colCnt="{len(ws)}" cellSpacing="0" borderFillIDRef="{borders((line, None))}" noAdjust="0">'
               f'<hp:sz width="{sum(ws)}" widthRelTo="ABSOLUTE" height="{rh * len(rows)}" heightRelTo="ABSOLUTE" protect="0"/>'
               '<hp:pos treatAsChar="0" affectLSpacing="0" flowWithText="1" allowOverlap="0" holdAnchorAndSO="0" vertRelTo="PARA" horzRelTo="COLUMN" '
               'vertAlign="TOP" horzAlign="LEFT" vertOffset="0" horzOffset="0"/><hp:outMargin left="0" right="0" top="0" bottom="0"/>'
               f'<hp:inMargin left="{pd}" right="{pd}" top="85" bottom="85"/>' + "".join(trs) + '</hp:tbl>')
        return para([((10, False, False, "000000"), tbl + "<hp:t/>")], "LEFT", 100)

    body = []
    for n, sh in enumerate(sheets):
        if n == 0:                                                 # 첫 문단에 쪽 설정(가로 A4·여백)을 단다
            body.append(para([((10, False, False, "000000"), _sec_pr() + t(""))], spacing=100))
        body.append(para([((16, True, False, "1F3864"), t(sh["title"]))], "CENTER", 160, page_break=n > 0))
        half = USABLE_MM / 2
        body.append(table([[(sh["company"], (9, True, False, "404040"), "left", None), (sh["printed"], (9, False, False, "404040"), "right", None)]],
                          [half, half], line=None, pad=0))
        for line in sh["info"]:
            if any(lab for _, lab in line):
                body.append(table([[(x, (9, lab, False, "000000"), "center" if lab else "left", "EDEDED" if lab else None) for x, lab in line]],
                                  [max(_text_mm(x, 9) + 2 * CELL_PAD_MM + 2, 22) for x, _ in line]))
            else:
                body.append(para([((9, False, True, "595959"), t("  ".join(x for x, _ in line)))]))
        body.append(para([((4, False, False, "000000"), t(""))], spacing=100))
        pt, mm = layout(sh)
        rows = [[(h, (pt, True, False, "000000"), "center", sh["header_fill"]) for h in sh["header"]]]
        rows += [[(d["text"], (pt, d["bold"], False, "000000"), d["align"], d["fill"]) for d in r] for r in sh["rows"]]
        body.append(table(rows, mm, repeat=True))
    section = XML_HEAD + f"<hs:sec {NS}>" + "".join(body) + "</hs:sec>"
    header = _header_xml(chars, paras, borders)               # 표·문단을 다 만든 뒤에 모양 목록이 완성된다
    title = escape(sheets[0]["title"]) if sheets else ""
    hpf = (XML_HEAD + f'<opf:package {NS} version="" unique-identifier="" id=""><opf:metadata><opf:title>{title}</opf:title><opf:language>ko</opf:language>'
           '<opf:meta name="creator" content="text">spec_to_production</opf:meta></opf:metadata><opf:manifest>'
           '<opf:item id="header" href="Contents/header.xml" media-type="application/xml"/><opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/>'
           '<opf:item id="settings" href="settings.xml" media-type="application/xml"/></opf:manifest>'
           '<opf:spine><opf:itemref idref="header" linear="yes"/><opf:itemref idref="section0" linear="yes"/></opf:spine></opf:package>')
    files = {
        "version.xml": XML_HEAD + '<hv:HCFVersion xmlns:hv="http://www.hancom.co.kr/hwpml/2011/version" tagetApplication="WORDPROCESSOR" major="5" minor="1" '
                                  'micro="1" buildNumber="0" os="1" xmlVersion="1.5" application="Hancom Office Hangul" appVersion="13, 0, 0, 1408 WIN32LEWindows_10"/>',
        "Contents/header.xml": header,
        "Contents/section0.xml": section,
        "Preview/PrvText.txt": "\r\n".join(f"{s['title']} {s['name']}" for s in sheets),
        "settings.xml": XML_HEAD + '<ha:HWPApplicationSetting xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app" xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0">'
                                   '<ha:CaretPosition listIDRef="0" paraIDRef="0" pos="0"/></ha:HWPApplicationSetting>',
        "META-INF/container.xml": XML_HEAD + '<ocf:container xmlns:ocf="urn:oasis:names:tc:opendocument:xmlns:container" xmlns:hpf="http://www.hancom.co.kr/schema/2011/hpf">'
                                             '<ocf:rootfiles><ocf:rootfile full-path="Contents/content.hpf" media-type="application/hwpml-package+xml"/>'
                                             '<ocf:rootfile full-path="Preview/PrvText.txt" media-type="text/plain"/></ocf:rootfiles></ocf:container>',
        "META-INF/manifest.xml": XML_HEAD + '<odf:manifest xmlns:odf="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>',
        "Contents/content.hpf": hpf,
    }
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)   # 맨 앞, 압축 없이
        for name, data in files.items():
            z.writestr(name, data.encode("utf-8"), compress_type=zipfile.ZIP_DEFLATED)


# ═════════════════════════ 6. 내보내기 + 형식 간 값 비교 ═════════════════════════
def export(xlsx_paths, formats=("pdf", "docx", "hwpx")):
    xlsx_paths = [Path(p) for p in xlsx_paths]
    for p in xlsx_paths:
        sheets = read_book(p)
        if "docx" in formats:
            to_docx(sheets, p.with_suffix(".docx"))
        if "hwpx" in formats:
            to_hwpx(sheets, p.with_suffix(".hwpx"))
    if "pdf" in formats:
        to_pdf(xlsx_paths)


def _norm(s):
    return re.sub(r"\s+", "", s)


def _norm_pdf(s):
    """PDF 글자 뽑기는 가운뎃점(·)을 빈칸으로 내기도 해서, PDF 비교에서는 빈칸처럼 뺀다."""
    return re.sub(r"[\s·]+", "", s)


def _docx_tables(path):
    from docx import Document
    return [[[c.text for c in row.cells] for row in t.rows] for t in Document(path).tables]


def _hwpx_tables(path):
    import xml.etree.ElementTree as ET
    hp = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("Contents/section0.xml"))
    return [[["".join(x.text or "" for x in tc.iter(hp + "t")) for tc in tr.findall(hp + "tc")] for tr in tbl.findall(hp + "tr")] for tbl in root.iter(hp + "tbl")]


def check_formats(xlsx, formats=("pdf", "docx", "hwpx")):
    """xlsx 와 같은 이름의 pdf·docx·hwpx 가 있으면, 표 칸 글자가 xlsx 와 같은지 본다. 틀린 곳을 글로 돌려준다."""
    xlsx = Path(xlsx)
    E = []
    sheets = read_book(xlsx)
    want = [[sh["header"]] + [[d["text"] for d in r] for r in sh["rows"]] for sh in sheets]
    heads = [x for sh in sheets for x in (sh["title"], sh["company"], sh["printed"])] + [x for sh in sheets for line in sh["info"] for x, _ in line]
    for ext, reader in ((".docx", _docx_tables), (".hwpx", _hwpx_tables)):
        if ext[1:] not in formats:
            continue
        p = xlsx.with_suffix(ext)
        if not p.exists():
            E.append(f"{p.name} 없음")
            continue
        tables = reader(p)
        got = [t for t in tables if t and any(t[0] == w[0] for w in want)]       # 머리글이 같은 표 = 본 표
        if len(got) != len(want):
            E.append(f"{p.name}: 본 표 {len(got)}개 (xlsx 시트 {len(want)}장)")
        for i, (g, w) in enumerate(zip(got, want)):
            if len(g) != len(w):
                E.append(f"{p.name} {sheets[i]['name']}: 줄 수 {len(g) - 1} ≠ xlsx {len(w) - 1}")
            bad = [(r, c, gv, wv) for r, (gr, wr) in enumerate(zip(g, w)) for c, (gv, wv) in enumerate(zip(gr, wr)) if gv != wv]
            for r, c, gv, wv in bad[:3]:
                E.append(f"{p.name} {sheets[i]['name']} {r}줄 {c + 1}칸: '{gv}' ≠ xlsx '{wv}'")
        flat = _norm(" ".join(c for t in tables for row in t for c in row))
        if ext == ".docx":
            from docx import Document
            flat += _norm(" ".join(p_.text for p_ in Document(p).paragraphs))
        else:
            with zipfile.ZipFile(p) as z:
                flat = _norm(re.sub(r"<[^>]+>", " ", z.read("Contents/section0.xml").decode("utf-8")).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">"))
        for h in heads:
            if _norm(h) not in flat:
                E.append(f"{p.name}: 머리 글 '{h}' 없음")
    p = xlsx.with_suffix(".pdf")
    if "pdf" not in formats:
        pass
    elif not p.exists():
        E.append(f"{p.name} 없음")
    else:
        import pymupdf
        with pymupdf.open(p) as d:
            text = _norm_pdf("".join(pg.get_text() for pg in d))
        cells = [c for w in want for row in w for c in row if c] + heads
        need = {}
        for c in cells:
            need[_norm_pdf(c)] = need.get(_norm_pdf(c), 0) + 1
        for c, k in need.items():
            if text.count(c) < k:                       # 같은 값이 여러 번이면 적어도 그만큼 나와야 한다
                E.append(f"{p.name}: '{c}' {k}번 중 {text.count(c)}번만 보임")
    return E
