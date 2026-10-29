"""제품 사양서(PDF)로 생산관리 가상 문서 11종(현업 모양 xlsx + 정답 csv)을 만든다.

입력: 한화비전 데이터시트(제품 사양서) PDF 1개 이상. 제품군(아날로그/네트워크/VPM)은 --category, 폴더 이름, 사양서 내용 순서로 정한다.
출력: 품목마스터·BOM·라우팅·작업장·비가동사유·MPS·생산계획·작업지시서·작업실적·비가동기록·4M변경점 폴더마다
      <문서>_원본.xlsx (ERP·MES 에서 내려받은 모양) 와 <문서>_정답.csv (같은 값의 깨끗한 표, basis 칸에 근거).
      EXPORT_DOCS(작업실적 빼고 10종)는 같은 xlsx 로 <문서>_원본.pdf/.docx/.hwpx 도 만든다 (production_export.py).
근거 표시: [사양서 p.N] 사양서에 그대로 있는 값 / [사양서 추론] 사양서 값에서 끌어낸 것 / [가상] 근거 없이 정한 값.
양식(칸·머리글)은 data/가상데이터/생산관리/ 의 기존 11종(공개 양식 조사로 정한 칸)을 따른다.
끝에 verify() 가 원본↔정답 줄 수, 문서 사이 키·수량·일정·날짜, 사양서 기준값, 그리고 PDF·docx·hwpx 표 칸이 xlsx 와 같은지 검사한다.

실행: python scripts/spec_to_production.py <사양서.pdf> [<사양서.pdf> ...] [--out 폴더] [--seed N] [--formats pdf,docx,hwpx] [--category network|analog|vpm]
필요: pip install pdfplumber openpyxl python-docx pymupdf / PDF 는 LibreOffice(soffice). 없으면 PDF 만 건너뛴다.
"""
import argparse
import csv
import importlib.util
import json
import logging
import math
import random
import re
import zlib
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path

try:
    import openpyxl
    import pdfplumber
except ImportError as e:                                   # 다른 PC 에서 처음 쓸 때 가장 흔한 문제
    raise SystemExit(f"필요한 파이썬 패키지가 없습니다: {e.name}\n  설치: pip install pdfplumber openpyxl python-docx pymupdf")
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from production_export import check_formats, export, find_soffice

logging.disable(logging.CRITICAL)
COMPANY = "(주)가상전자 · 생산1공장"
TODAY = datetime(2026, 11, 12, 17, 0)        # 출력 기준 시각: 이때까지 일어난 일만 실적·기록에 남는다
FIRST_MONDAY = date(2026, 11, 2)
WEEKS = 8
LOT_MAX = 200
WORK = [(8 * 60, 12 * 60), (13 * 60, 17 * 60)]
WEEKDAY = "월화수목금토일"
OUT_FORMATS = ("pdf", "docx", "hwpx")
FOLDERS = ["품목마스터", "BOM", "라우팅", "작업장", "비가동사유", "MPS", "생산계획", "작업지시서", "작업실적", "비가동기록", "4M변경점"]
EXPORT_DOCS = [f for f in FOLDERS if f != "작업실적"]   # xlsx 말고 PDF·docx·hwpx 로도 내는 문서. 작업실적은 머리글 2줄·병합 칸 일지 양식이라 아직 못 낸다
CATEGORY_BY_DIR = {"네트워크 카메라": "network", "아날로그 카메라": "analog", "VPM": "vpm"}
CATEGORY_NAME = {"network": "네트워크 카메라", "analog": "아날로그 카메라", "vpm": "VPM 카메라"}
THIN = Side(style="thin", color="999999")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


# ═════════════════════════ 1. 사양서 읽기 ═════════════════════════
def _num(s):
    return float(s.replace(",", ""))


# 파일 이름의 '_TND-F6022R,' 와 붙어 쓴 'TNU-L2320ETNU-L2120E' 도 모델별로 끊어 읽는다
MODEL_RE = r"(?<![a-z0-9-])[A-Z]{2,4}-[A-Z]{0,2}\d{3,5}[A-Z0-9]*?(?=[A-Z]{3}-|[^A-Za-z0-9]|$)"


def read_specs(pdf, category=None):
    """PDF 한 개에서 제품 사양을 모델마다 하나씩 읽는다. 브로셔처럼 한 표에 모델 여러 개가 열로 나란히 있으면
    머리 줄의 모델 순서대로 열 값을 나눠 갖는다 (예: 'IR LED 45m 30m' → 첫 모델 45m, 둘째 모델 30m)."""
    pdf = Path(pdf)
    if re.match(r"DataSheet_", pdf.name):
        return [read_spec(pdf, category)]
    with pdfplumber.open(pdf) as d:
        full = "\n".join((pg.extract_text() or "") for pg in d.pages)
    models = list(dict.fromkeys(re.findall(MODEL_RE, pdf.stem))) or list(dict.fromkeys(re.findall(MODEL_RE, full)))[:1]
    if len(models) > 1:
        for l in full.splitlines():
            found = list(dict.fromkeys(re.findall(MODEL_RE, l)))
            if set(models) <= set(found):
                models = found
                break
    return [read_spec(pdf, category, model=mdl, col=i, ncol=len(models)) for i, mdl in enumerate(models)]


def read_spec(pdf, category=None, model=None, col=0, ncol=1):
    pdf = Path(pdf)

    def pick(vals):
        """모델별 열 값이 모델 수만큼 있으면 이 모델 열의 값, 아니면 첫 값."""
        return vals[col] if ncol > 1 and len(vals) == ncol else vals[0]
    with pdfplumber.open(pdf) as d:
        pages = [(pg.extract_text() or "") for pg in d.pages]
    full = "\n".join(pages)
    lines = [l.strip() for l in full.splitlines()]

    def page_of(text):
        for i, t in enumerate(pages, 1):
            if text and text in t:
                return i
        return None

    def line_val(label_pat):
        """'라벨 값' 한 줄에서 값과 쪽 번호를 꺼낸다."""
        for l in lines:
            m = re.match(label_pat, l)
            if m:
                return {"v": m.group(1).strip(), "p": page_of(l)}
        return None

    m = re.match(r"DataSheet_(.+?)_\d", pdf.name)
    model = model or (m.group(1) if m else (re.findall(MODEL_RE, full) or [pdf.stem])[0])
    if category is not None:
        category_basis = "직접 지정(--category)"
    elif hit := [p for p in pdf.parts if p in CATEGORY_BY_DIR]:
        category, category_basis = CATEGORY_BY_DIR[hit[0]], f"폴더 이름 '{hit[0]}'"
    else:   # 폴더 이름이 없으면 사양서 글로 정한다 (카메라 사양서 178개로 시험: 아날로그 17/17, VPM 15/17)
        ip = re.search(r"PoE|이더넷|RJ-45|ONVIF", full)
        if re.search(r"\bVPM\b", full):
            category = "vpm"
        elif re.search(r"AHD|TVI|CVI|CVBS", full) and not ip:
            category = "analog"
        else:
            category = "network"
        category_basis = "사양서 내용으로 추정"
    first = [l.strip() for l in pages[0].splitlines()[:8]]
    ok_line = [l for l in first if model not in l and "Hanwha" not in l and not l.startswith("•")]
    title = next((l for l in ok_line if l.endswith("카메라")), "") or next((l for l in ok_line if "카메라" in l or "센서" in l), "")
    if not title:
        title = next((l for l in first[1:6] if l and l != model and l != "'" and not l.startswith(model)), CATEGORY_NAME[category])
    s = {"model": model, "title": title, "category": category, "category_basis": category_basis, "pages": len(pages), "file": pdf.name}

    # 무게 (g, kg, Kg) — '제품 무게 1935g', '제품 무게 4.53kg', '제품 사이즈 / 무게 157x111x482mm, 2050g'
    s["weight_g"] = None
    for i, l in enumerate(lines):
        if "무게" in l and not l.startswith("•"):
            ws = re.findall(r"([\d,]+(?:\.\d+)?)\s*(kg|Kg|KG|g)\b", l)
            if ws:
                nxt = lines[i + 1] if i + 1 < len(lines) else ""       # 'Camera: 1897g LED: 996g Total:' 다음 줄에 합계가 오는 경우
                tot = re.search(r"Total:\s*([\d,]+(?:\.\d+)?)\s*(kg|Kg|KG|g)", l + " " + nxt)
                v, u = tot.groups() if tot else (ws[0] if "제품 무게" in l else ws[-1])
                s["weight_g"] = {"v": round(_num(v) * (1000 if u.lower() == "kg" else 1), 1), "p": page_of(l)}
                break
    s["dims"] = line_val(r"^제품 치수\s+(.+)$") or line_val(r"^제품 사이즈 / 무게\s+(.+?),\s*[\d,.]+\s*(?:kg|g)")
    s["color"] = line_val(r"^색상\s+(?!/)(.+)$")
    s["material"] = line_val(r"^재질\s+(?!관리)(.+)$")
    cm = line_val(r"^색상 / 재질\s+(.+)$")
    if cm and not s["color"]:
        a, b = (cm["v"].split(" / ", 1) if " / " in cm["v"] else (cm["v"].split(", ", 1) + [""])[:2])
        s["color"] = {"v": a.strip(), "p": cm["p"]}
        s["material"] = {"v": b.strip(), "p": cm["p"]} if b.strip() else None
    s["ral"] = line_val(r"^RAL 코드\s+(\S+)")
    sz = re.search(r'촬상 소자(?: 크기)?\s+(?:Diagonal [\d.]+mm\s+)?(1/[\d.]+|\d(?:\.\d+)?(?="))"?', full)
    s["sensor"] = {"v": sz.group(1) + '"', "p": page_of(sz.group(0))} if sz else None
    bolo = re.search(r"비냉각식 마이크로 볼로미터\s*\((\d+\s*x\s*\d+)\)", full)
    s["thermal"] = {"v": bolo.group(1).replace(" ", ""), "p": page_of(bolo.group(0))} if bolo else None
    netd = re.search(r"NETD\s*([<≤]?\s*\d+\s*mK)", full)
    s["netd"] = {"v": netd.group(1).replace(" ", ""), "p": page_of(netd.group(0))} if netd else None
    res = re.search(r"^해상도\s+(\d+\s*[x×]\s*\d+)", full, re.M) or re.search(r"스트림\s*1\s*:\s*(\d+\s*[x×]\s*\d+)", full)
    s["resolution"] = {"v": res.group(1).replace(" ", "").replace("×", "x"), "p": page_of(res.group(0))} if res else None
    s["ram"] = line_val(r"^메모리\s+RAM\s+(\S+)")
    s["flash"] = line_val(r"^Flash\s+(.+)$")
    s["edge"] = line_val(r"^Edge Storage\s+(.+)$")
    s["ethernet"] = line_val(r"^이더넷\s+(.+)$")
    s["video_out"] = line_val(r"^영상 출력\s+(.+)$")
    s["alarm"] = line_val(r"^알람 입출력\s+입/출력\s+(.+)$") or line_val(r"^알람 입출력\s+(.+)$")
    s["audio_in"] = line_val(r"^오디오 입력\s+(.+)$")
    s["focal"] = line_val(r"^초점거리(?: \(mm\))?\s+([\d.]+(?:\s*[~-]\s*[\d.]+)?(?:mm)?(?:\s+[\d.]+(?:\s*[~-]\s*[\d.]+)?)?)")
    if s["focal"]:
        s["focal"]["v"] = s["focal"]["v"].replace("mm", "").replace("-", "~")
    s["fno"] = line_val(r"^최대 구경비 \(F number\)\s+광각\s+(.+)$")     # Bi-spectrum 은 '1.0 1.7' 처럼 열화상·실화상 두 값
    s["zoom"] = line_val(r"^줌 배율\s+광학\s+([\d.]+)")
    s["pan_speed"] = line_val(r"^팬 속도\s+(?:프리셋\s+)?(.+)$")
    s["preset_acc"] = line_val(r"^프리셋 정확도\s+(.+)$")
    s["ir_nm"] = line_val(r"^IR 파장\s+(\d+\s*nm)")
    ir = re.search(r"(?:IR 가시거리|조명 가시거리 IR|야간 가시거리|IR\s+WiseIR)\D{0,12}?(\d+)\s*m", full)
    s["ir_m"] = {"v": int(ir.group(1)), "p": page_of(ir.group(0))} if ir else None
    irl = re.search(r"^IR LED\s+(\d+)\s*m((?:\s+\d+\s*m)*)", full, re.M)          # 브로셔: 'IR LED 45m 30m'
    if irl and not s["ir_m"]:
        s["ir_m"] = {"v": int(pick([irl.group(1)] + re.findall(r"(\d+)\s*m", irl.group(2)))), "p": page_of(irl.group(0))}
    fd = re.search(r"감지 거리 / 각도\s+([\d.]+\s*m\s*/\s*[\d.]+\s*°)", full)
    s["flame"] = {"v": fd.group(1).replace(" ", ""), "p": page_of(fd.group(0))} if fd else None
    fn = re.search(r"감지센서\s+((?:\d+개\s*)+)", full)
    s["flame_n"] = {"v": int(pick(re.findall(r"(\d+)개", fn.group(1)))), "p": page_of(fn.group(0))} if fn else None
    mk = re.search(r"제조 및 A/S\s*:\s*(\S+)", full)
    s["maker"] = {"v": mk.group(1), "p": page_of(mk.group(0))} if mk else None
    s["power_in"] = line_val(r"^전원\s+(.+)$")
    # 소비 전력: '소비 전력' 줄과 이어지는 전원별 줄에서 전원별 최대값
    s["power"] = []
    for i, l in enumerate(lines):
        if re.match(r"^소비\s?전력", l):
            seg = [re.sub(r"^소비\s?전력\s*", "", l)]
            for nx in lines[i + 1:i + 4]:
                if re.match(r"^(H?PoE|Hi-PoE|UPoE|PoE\+|\d+\s?V\s?(AC|DC)|AC\d+|DC\d+)", nx):
                    seg.append(nx)
                else:
                    break
            for g in seg:
                src = re.match(r"^(H?PoE\+*|Hi-PoE|UPoE|\d+\s?V\s?(?:AC|DC)|AC\d+|DC\d+)", g)
                mx = re.search(r"(?:max|최대)\.?\s*:?\s*([\d.]+)\s*W", g, re.I)
                nums = re.findall(r"([\d.]+)\s*W", g)
                v = mx.group(1) if mx else (pick(nums) if nums else None)      # 브로셔: '7W 6W' 는 모델별 열 값
                if v:
                    s["power"].append({"src": src.group(1) if src else "전원", "v": float(v), "p": page_of(l)})
            break
    # 동작 온도: 여러 범위가 적혀 있으면 모두 만족하는 구간
    s["op_temp"] = None
    for i, l in enumerate(lines):
        if re.match(r"^동작 (조건 )?온도", l):
            txt = l + " " + (lines[i + 1] if i + 1 < len(lines) and re.match(r"^(IR|AC|HPoE|PoE)", lines[i + 1]) else "")
            rs = re.findall(r"(-?\d+)\s*[°℃]\s*C?\s*~\s*\+?(-?\d+)\s*[°℃]", txt)
            if rs and ncol > 1 and len(rs) == ncol:                    # 브로셔: 모델별 열 값
                a, b = rs[col]
                s["op_temp"] = {"v": (int(a), int(b)), "p": page_of(l)}
            elif rs:
                s["op_temp"] = {"v": (max(int(a) for a, b in rs), min(int(b) for a, b in rs)), "p": page_of(l)}
            break
    s["ip_list"] = sorted(set(re.findall(r"IP6K9K|IP\d{2}", full)))
    s["ik_list"] = sorted(set(re.findall(r"IK\d{2}\+?", full)))
    ch = re.search(r"(\d)\s*채널", title)
    s["channels"] = int(ch.group(1)) if ch and int(ch.group(1)) > 1 else (2 if "듀얼 센서" in title else 1)

    def has(pat):
        mm = re.search(pat, full, re.M)
        return {"v": True, "p": page_of(mm.group(0))} if mm else None
    s["flags"] = {
        "poe": has(r"PoE"), "ethernet": has(r"^이더넷"), "sfp": has(r"SFP"), "m12": has(r"M12"),
        "bnc": has(r"BNC"), "rs485": has(r"RS-485"), "alarm": has(r"^알람 입출력"), "audio": has(r"^오디오 (입력|출력)"),
        "mic": has(r"내장형? 마이크"), "speaker": has(r"내장형? 스피커|스피커 출력"), "call": has(r"호출"),
        "sd": has(r"Micro ?SD|SD/SDHC"), "usb": has(r"USB:"), "tpm": has(r"TPM|시큐어 부트|펌웨어 암호화"),
        "icr": has(r"Day & Night"), "ir": has(r"IR LED|IR 파장|조명 종류\s+IR"), "mount": has(r"마운트 타입"),
        "motor_pt": has(r"^팬 속도|시퀀스 프리셋"), "slip": has(r"팬\s+0°~360°|360° Endless|Endless"), "wiper": has(r"와이퍼|Wiper"),
        "heater": has(r"히터|Heater"), "vent": has(r"AIR vent|에어벤트"), "exproof": has(r"Ex db|방폭"),
        "anti_ligature": has(r"결박 방지"), "ip": has(r"IP6\d|IP6K9K"),
    }
    return s


def basis(spec, key):
    v = spec.get(key) if key in spec else spec["flags"].get(key)
    if isinstance(v, dict) and v.get("p"):
        return f"[사양서 p.{v['p']}]"
    if isinstance(v, list) and v and isinstance(v[0], dict) and v[0].get("p"):
        return f"[사양서 p.{v[0]['p']}]"
    return "[사양서 추론]"


def sv(spec, key, default=""):
    v = spec.get(key)
    return v["v"] if isinstance(v, dict) else default


def fl(spec, key):
    return bool(spec["flags"].get(key))


# ═════════════════════════ 2. 근무 달력 (평일 08–12, 13–17) ═════════════════════════
def next_work(t):
    while True:
        d = t.date()
        if d.weekday() >= 5:
            t = datetime.combine(d + timedelta(days=1), time(8))
            continue
        m = t.hour * 60 + t.minute
        for a, b in WORK:
            if m < a:
                return datetime.combine(d, time(0)) + timedelta(minutes=a)
            if a <= m < b:
                return t
        t = datetime.combine(d + timedelta(days=1), time(8))


def add_work(t, minutes):
    t = next_work(t)
    rem = minutes
    while rem > 0:
        m = t.hour * 60 + t.minute
        seg_end = next(b for a, b in WORK if a <= m < b)
        if rem <= seg_end - m:
            return t + timedelta(minutes=rem)
        rem -= seg_end - m
        t = next_work(datetime.combine(t.date(), time(0)) + timedelta(minutes=seg_end))
    return t


def segments(s, e):
    out, cur = [], next_work(s)
    while cur < e:
        m = cur.hour * 60 + cur.minute
        seg_end = datetime.combine(cur.date(), time(0)) + timedelta(minutes=next(b for a, b in WORK if a <= m < b))
        out.append((cur, min(e, seg_end)))
        cur = next_work(seg_end)
    return out


def wmin(s, e):
    return sum(round((b - a).total_seconds() / 60) for a, b in segments(s, e))


def workdays(d0, d1):
    d, out = d0, []
    while d <= d1:
        if d.weekday() < 5:
            out.append(d)
        d += timedelta(days=1)
    return out


def ds(d):
    return d.strftime("%Y-%m-%d")


def dts(t):
    return t.strftime("%Y-%m-%d %H:%M")


# ═════════════════════════ 3. 제품 설계 (사양서 → 반제품·부품·공정) ═════════════════════════
def rm_code(name, spec_text):
    return "RM-%05d" % (zlib.crc32((name + "|" + spec_text).encode()) % 100000)


def design(spec, rng):
    """반제품, 부품 줄, 공정을 사양서 특징에서 정한다.
    부품 줄: (반제품키 또는 None=완제품 직속, 품명, 규격, 수량, 단위, 손실률, 투입공정, 근거)."""
    res = sv(spec, "resolution")
    net = fl(spec, "ethernet")
    n_ch = spec["channels"]
    semis = [("MB", "메인 PCB 조립체", "SMT", 1, "[사양서 추론] 영상 처리 기판")]
    if spec["sensor"] or not spec["thermal"]:
        semis.append(("OP", "광학 모듈", "광학 조립", n_ch,
                      f"{basis(spec, 'sensor')} 촬상 소자" + (f" / 제품명의 {n_ch}채널" if n_ch > 1 else "")))
    if spec["thermal"]:
        semis.append(("TH", "열화상 모듈", "광학 조립", 1, f"{basis(spec, 'thermal')} 비냉각식 마이크로 볼로미터"))
    if fl(spec, "ir"):
        semis.append(("IR", "IR 조명 보드", "SMT", 1, f"{basis(spec, 'ir')} IR 조명"))
    if fl(spec, "motor_pt"):
        semis.append(("PT", "팬/틸트 구동부", "팬/틸트 조립", 1, f"{basis(spec, 'motor_pt')} 전동 팬/틸트 (팬 속도·프리셋 사양)"))
    semis.append(("HS", "하우징 조립체", "하우징 조립", 1, f"{basis(spec, 'material')} 외관 재질" if spec["material"] else "[사양서 추론] 외관"))

    P = [("MB", "메인 PCB", "FR-4 6층", 1, "EA", 0.5, "SMT", "[가상] 층수·재질"),
         ("MB", "영상처리 SoC", f"{res} 영상 처리" if res else "영상 처리", 1, "EA", 0.2, "SMT",
          f"{basis(spec, 'resolution')} 해상도 {res} / 부품 선정 [가상]" if res else "[사양서 추론] / 부품 선정 [가상]"),
         ("MB", "전원 IC", "DC-DC 컨버터", 1, "EA", 0.3, "SMT", "[가상]")]
    if spec["ram"]:
        P.append(("MB", "DRAM", f"RAM {sv(spec, 'ram')}", 1, "EA", 0.2, "SMT", f"{basis(spec, 'ram')} 메모리 RAM {sv(spec, 'ram')}"))
    elif net:
        P.append(("MB", "DRAM", "", 1, "EA", 0.2, "SMT", "[사양서 추론] 네트워크 영상 처리용"))
    P.append(("MB", "플래시 메모리", sv(spec, "flash"), 1, "EA", 0.2, "SMT",
              f"{basis(spec, 'flash')} Flash {sv(spec, 'flash')}" if spec["flash"] else "[사양서 추론] 펌웨어 저장"))
    if fl(spec, "tpm"):
        P.append(("MB", "TPM 보안 칩", "", 1, "EA", 0.1, "SMT", f"{basis(spec, 'tpm')} 보안(시큐어 부트·펌웨어 보호)"))
    if net:
        eth = sv(spec, "ethernet")
        speed = re.search(r"(10/100(?:/1000)?BASE-T)", eth)
        if fl(spec, "m12"):
            P.append(("MB", "M12 이더넷 커넥터", speed.group(1) if speed else "", 1, "EA", 0.3, "하우징 조립", f"{basis(spec, 'm12')} M12 커넥터"))
        else:
            P.append(("MB", "이더넷 커넥터 RJ-45", ((speed.group(1) if speed else "") + (", 금속 차폐" if "shielded" in eth else "")).strip(", "), 1, "EA", 0.3, "SMT",
                      f"{basis(spec, 'ethernet')} 이더넷 {eth}"))
        if fl(spec, "sfp"):
            P.append(("MB", "SFP 케이지", "", 1, "EA", 0.3, "SMT", f"{basis(spec, 'sfp')} SFP"))
    pin = sv(spec, "power_in")
    if fl(spec, "poe"):
        poe = re.search(r"((?:H|Hi-|U)?PoE\+*(?:\([^)]*\))?)", pin) if pin else None
        P.append(("MB", "PoE 전원 모듈", poe.group(1) if poe else "PoE", 1, "EA", 0.3, "SMT",
                  f"{basis(spec, 'power_in')} 전원 {pin}" if pin else f"{basis(spec, 'poe')} PoE"))
    dc = re.findall(r"\d+\s?V\s?(?:AC|DC)|(?:AC|DC)\s?\d+\s?V", pin) if pin else []
    if dc:
        P.append(("MB", "전원 입력 단자", ", ".join(dc), 1, "EA", 0.3, "하우징 조립", f"{basis(spec, 'power_in')} 전원 {pin}"))
    if fl(spec, "bnc"):
        vo = sv(spec, "video_out")
        P.append(("MB", "BNC 커넥터", re.sub(r"^BNC:\s*", "", vo) if "BNC" in vo else "", 1, "EA", 0.3, "하우징 조립", f"{basis(spec, 'bnc')} 영상 출력 BNC"))
    if fl(spec, "rs485"):
        P.append(("MB", "RS-485 트랜시버", "", 1, "EA", 0.2, "SMT", f"{basis(spec, 'rs485')} RS-485"))
    if fl(spec, "alarm"):
        P.append(("MB", "알람 I/O 단자", sv(spec, "alarm")[:30], 1, "EA", 0.3, "하우징 조립", f"{basis(spec, 'alarm')} 알람 입출력"))
    if fl(spec, "audio"):
        P.append(("MB", "오디오 코덱", "", 1, "EA", 0.2, "SMT", f"{basis(spec, 'audio')} 오디오 입출력"))
    if fl(spec, "mic"):
        P.append(("MB", "내장 마이크", "", 1, "EA", 0.5, "하우징 조립", f"{basis(spec, 'mic')} 내장 마이크"))
    if fl(spec, "sd"):
        edge = sv(spec, "edge")
        slots = re.search(r"(\d)\s*슬롯", edge)
        P.append(("MB", "microSD 슬롯", re.split(r"\s\*", edge)[0].strip(), int(slots.group(1)) if slots else 1, "EA", 0.3, "SMT",
                  f"{basis(spec, 'edge')} Edge Storage" if edge else f"{basis(spec, 'sd')} SD"))
    if fl(spec, "usb"):
        P.append(("MB", "Micro USB 커넥터", "Type B", 1, "EA", 0.3, "SMT", f"{basis(spec, 'usb')} USB 출력"))
    focal_t = sv(spec, "focal").split()
    fno_t = re.findall(r"[\d.]+", sv(spec, "fno"))
    bi = bool(spec["thermal"]) and bool(spec["sensor"])          # Bi-spectrum: 첫 값 열화상, 둘째 값 실화상
    vis_focal = (focal_t[1] if bi and len(focal_t) > 1 else focal_t[0]) if focal_t else ""
    vis_fno = (fno_t[1] if bi and len(fno_t) > 1 else fno_t[0]) if fno_t else ""
    th_focal = focal_t[0] if focal_t and spec["thermal"] else ""
    th_fno = fno_t[0] if fno_t and spec["thermal"] else ""
    if any(k == "OP" for k, *_ in semis):
        size = sv(spec, "sensor")
        P.append(("OP", "이미지 센서", f"{size} CMOS".strip(), 1, "EA", 0.3, "광학 조립", f"{basis(spec, 'sensor')} 촬상 소자 {size}".strip()))
        zoom = sv(spec, "zoom")
        if fl(spec, "mount"):
            P.append(("OP", "렌즈 마운트 링", "C/CS 마운트", 1, "EA", 0.3, "광학 조립", f"{basis(spec, 'mount')} 마운트 타입 (렌즈는 별도 구매)"))
        else:
            lspec = " ".join(x for x in [f"{vis_focal}mm" if vis_focal else "", f"광학 {zoom}배" if zoom else "", f"F{vis_fno}" if vis_fno else ""] if x)
            P.append(("OP", "전동 줌 렌즈" if zoom else "고정 초점 렌즈", lspec, 1, "EA", 0.5, "광학 조립",
                      f"{basis(spec, 'focal')} 초점거리·구경비" if vis_focal else "[사양서 추론] 렌즈"))
        if fl(spec, "icr"):
            P.append(("OP", "ICR 필터", "", 1, "EA", 0.3, "광학 조립", f"{basis(spec, 'icr')} Day & Night(ICR)"))
        P.append(("OP", "센서 보드", "", 1, "EA", 0.3, "광학 조립", "[사양서 추론]"))
    if spec.get("flame_n"):
        P.append(("OP" if any(k == "OP" for k, *_ in semis) else "MB", "불꽃감지 IR 센서", f"감지 {sv(spec, 'flame')}" if spec.get("flame") else "",
                  sv(spec, "flame_n"), "EA", 0.5, "광학 조립" if any(k == "OP" for k, *_ in semis) else "SMT", f"{basis(spec, 'flame_n')} 불꽃 감지센서 개수"))
    if spec["thermal"]:
        P.append(("TH", "열화상 검출기", f"비냉각식 마이크로 볼로미터 {sv(spec, 'thermal')}" + (f", NETD {sv(spec, 'netd')}" if spec["netd"] else ""), 1, "EA", 0.2,
                  "광학 조립", f"{basis(spec, 'thermal')} 열화상 소자"))
        P.append(("TH", "열화상 렌즈", " ".join(x for x in [f"{th_focal}mm" if th_focal else "", f"F{th_fno}" if th_fno else ""] if x), 1, "EA", 0.3, "광학 조립",
                  f"{basis(spec, 'focal')} 열화상 초점거리·구경비" if th_focal else "[사양서 추론] 열화상 광학계"))
    if fl(spec, "ir"):
        rng_m = sv(spec, "ir_m", 0) or 0
        n_led = 4 if rng_m <= 20 else 8 if rng_m <= 50 else 12 if rng_m <= 100 else 16
        P.append(("IR", "IR LED", sv(spec, "ir_nm").replace(" ", ""), n_led, "EA", 1.0, "SMT",
                  (f"{basis(spec, 'ir_nm')} IR 파장" if spec["ir_nm"] else "[사양서 추론] IR") + (f" / 개수는 가시거리 {rng_m}m 기준 [가상]" if rng_m else " / 개수 [가상]")))
        P.append(("IR", "IR 보드 PCB", "FR-4 2층", 1, "EA", 0.5, "SMT", "[가상] 층수"))
    if fl(spec, "motor_pt"):
        P += [("PT", "팬 모터", sv(spec, "pan_speed")[:30], 1, "EA", 0.3, "팬/틸트 조립", f"{basis(spec, 'pan_speed')} 팬 속도" if spec["pan_speed"] else f"{basis(spec, 'motor_pt')} 전동 팬"),
              ("PT", "틸트 모터", "", 1, "EA", 0.3, "팬/틸트 조립", "[사양서 추론] 틸트 구동"),
              ("PT", "팬/틸트 구동 보드", "", 1, "EA", 0.3, "팬/틸트 조립", "[사양서 추론]")]
        if fl(spec, "slip"):
            P.append(("PT", "슬립링", "360° 연속 회전", 1, "EA", 0.3, "팬/틸트 조립", f"{basis(spec, 'slip')} 팬 360°"))
    if fl(spec, "speaker"):
        P.append(("HS", "내장 스피커", "", 1, "EA", 0.5, "하우징 조립", f"{basis(spec, 'speaker')} 내장 스피커"))
    if fl(spec, "call"):
        P.append(("HS", "호출 센서", "비접촉", 1, "EA", 0.5, "하우징 조립", f"{basis(spec, 'call')} 호출 기능"))
    if fl(spec, "wiper"):
        P.append(("HS", "와이퍼 모터", "", 1, "EA", 0.3, "하우징 조립", f"{basis(spec, 'wiper')} 와이퍼"))
    if fl(spec, "heater"):
        P.append(("HS", "히터", "", 1, "EA", 0.3, "하우징 조립", f"{basis(spec, 'heater')} 히터"))
    mat, col, ral = sv(spec, "material"), sv(spec, "color"), sv(spec, "ral")
    hs_spec = ", ".join(x for x in [mat, col, ral, "내압방폭(Ex db)" if fl(spec, "exproof") else ""] if x)
    P.append(("HS", "하우징 본체", hs_spec, 1, "EA", 0.5, "하우징 조립", f"{basis(spec, 'material')} 재질·색상" if mat else "[사양서 추론]"))
    P.append(("HS", "돔 커버" if ("돔" in mat or "돔" in spec["title"]) else "전면 윈도우", "", 1, "EA", 0.5, "하우징 조립", "[사양서 추론] 렌즈 보호창"))
    if spec["ip_list"]:
        P.append(("HS", "방수 개스킷", ", ".join(spec["ip_list"]), 1, "EA", 1.0, "하우징 조립", f"{basis(spec, 'ip')} 방진·방수 {', '.join(spec['ip_list'])}"))
    if fl(spec, "vent"):
        P.append(("HS", "에어벤트", "습도 제어", 1, "EA", 0.5, "하우징 조립", f"{basis(spec, 'vent')} AIR vent"))
    if fl(spec, "exproof"):
        P.append(("HS", "방폭 케이블 글랜드", "Ex d", 1, "EA", 0.5, "하우징 조립", f"{basis(spec, 'exproof')} 방폭 인증"))
    P.append(("HS", "특수 나사(결박 방지)" if fl(spec, "anti_ligature") else "체결 나사", "M3x8 SUS", rng.choice([4, 6, 8]), "EA", 2.0, "하우징 조립",
              f"{basis(spec, 'anti_ligature')} 결박 방지 설계 / 규격·개수 [가상]" if fl(spec, "anti_ligature") else "[가상] 규격·개수"))
    P += [(None, "포장 박스", "골판지 이중", 1, "EA", 0.5, "포장", "[가상]"),
          (None, "설치 설명서", "다국어", 1, "EA", 0.5, "포장", "[가상]"),
          (None, "제품 라벨", "시리얼·MAC" if net else "시리얼", 1, "EA", 1.0, "포장", "[사양서 추론] 시리얼 관리 / 라벨 규격 [가상]")]

    # 공정 (코드, 공정명, 작업장, 기본 표준시간(분/EA), 준비시간, 인원, 검사 항목, 근거)
    wkg = (sv(spec, "weight_g", 0) or 0) / 1000
    k = 1 + len(P) / 40 + min(wkg, 40) / 20
    ops = [("OP-SMT", "SMT", "SMT라인", 1.2, 30, 2, "부품 실장 상태(AOI)", "[사양서 추론] 공정 / 시간·인원 [가상]")]
    if net or fl(spec, "tpm"):
        ops.append(("OP-FW", "FW 기록", "펌웨어기록실", 0.8, 10, 1, "FW 버전 일치" + (", 시큐어 부트·기기인증서 기록" if fl(spec, "tpm") else ""),
                    f"{basis(spec, 'tpm') if fl(spec, 'tpm') else basis(spec, 'ethernet')} / 시간 [가상]"))
    oi = ["초점·화각"]
    if fl(spec, "icr"):
        oi.append("ICR 전환")
    if spec["zoom"]:
        oi.append(f"줌 {sv(spec, 'zoom')}배 동작")
    if spec["netd"]:
        oi.append(f"NETD {sv(spec, 'netd')}")
    ops.append(("OP-OPT", "광학 조립", "광학조립실", 2.5 * n_ch, 15, 2, ", ".join(oi),
                f"{basis(spec, 'focal') if spec['focal'] else basis(spec, 'sensor')} 광학 사양 / 시간 [가상]"))
    if fl(spec, "motor_pt"):
        ops.append(("OP-PT", "팬/틸트 조립", "조립1라인", 5.0, 20, 2, "구동 소음·걸림 없음", f"{basis(spec, 'motor_pt')} 전동 팬·틸트 / 시간 [가상]"))
    ops.append(("OP-ASM", "하우징 조립", "조립1라인", 4.0, 10, 3, "체결 토크, 외관" + (", 결박 방지 나사" if fl(spec, "anti_ligature") else ""), "[사양서 추론] 하우징 / 시간 [가상]"))
    if fl(spec, "exproof"):
        ops.append(("OP-EXP", "내압 시험(방폭)", "방폭시험실", 3.0, 20, 1, "내압 용기 압력 시험", f"{basis(spec, 'exproof')} 방폭 인증 / 시험 조건 [가상]"))
    if spec["ip_list"]:
        ops.append(("OP-SEAL", "기밀 시험", "기밀시험기", 1.8, 20, 1, f"누설 없음({', '.join(spec['ip_list'])})", f"{basis(spec, 'ip')} 방진·방수 등급 / 시험 조건 [가상]"))
    fi = ["BNC 영상 출력" if fl(spec, "bnc") else "영상 스트림"]
    for key, label in (("ir", "IR 점등"), ("audio", "오디오 입출력"), ("alarm", "알람 I/O"), ("sd", "SD 인식"), ("poe", "PoE 동작"), ("rs485", "RS-485 통신"),
                       ("speaker", "스피커 출력"), ("call", "호출 동작")):
        if fl(spec, key):
            fi.append(label)
    if spec.get("flame"):
        fi.append(f"불꽃 감지 동작({sv(spec, 'flame')})")
    if fl(spec, "motor_pt"):
        fi.append("팬/틸트 동작" + (f"(프리셋 정확도 {sv(spec, 'preset_acc')})" if spec["preset_acc"] else ""))
    pw = spec["power"][0] if spec["power"] else None
    if pw:
        fi.append(f"소비전력 {pw['src']} {pw['v']:g}W 이하")
    ops.append(("OP-FCT", "기능 검사", "검사라인", 3.5, 15, 2, ", ".join(fi), (f"[사양서 p.{pw['p']}] 소비전력 / " if pw else "") + "검사 항목은 사양 기능 [사양서 추론]"))
    t = spec["op_temp"]
    ops.append(("OP-AGE", "에이징", "에이징실", 0.5, 30, 1, f"동작 온도({t['v'][0]}~{t['v'][1]}°C) 내 연속 동작" if t else "상온 연속 동작",
                f"[사양서 p.{t['p']}] 동작 조건 / 시험 시간 [가상]" if t else "[가상]"))
    wt = sv(spec, "weight_g", 0)
    ops.append(("OP-PKG", "포장", "포장라인", 1.5, 5, 1, f"중량 {wt:g}g ±3%, 모델·라벨 일치" if wt else "모델·라벨 일치",
                f"{basis(spec, 'weight_g')} 제품 무게 / 허용치 [가상]" if wt else "[가상]"))
    routing = [dict(code=c, name=n, wc=w, std=round(b * k, 2), setup=s_, crew=cr, insp=i, basis=bs) for c, n, w, b, s_, cr, i, bs in ops]
    return {"semis": semis, "parts": P, "routing": routing}


# ═════════════════════════ 4. 엑셀 쓰기 도우미 ═════════════════════════
def head(ws, title, right="", sub="", size=15):
    """1행 제목, 2행 회사·출력일, 3행 조회조건, 4행 빈 줄 (기존 양식과 같은 머리)."""
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=size)
    ws["A2"] = COMPANY
    if right:
        ws["H2"] = right
    if sub:
        ws["A3"] = sub
    ws.append([])


def title_block(ws, ncols, merge_to=None):
    """표 위쪽 3줄을 보기 좋게 꾸민다: 제목은 표 너비로 가운데, 출력일은 2행 오른쪽 끝, 조회조건은 3행 왼쪽."""
    last = merge_to or ncols
    if ncols >= 2:
        for c in list(ws[2]):                                    # 'H2' 에 둔 출력일을 표 오른쪽 끝으로
            if isinstance(c.value, str) and c.value.startswith("출력일") and c.column != ncols:
                ws.cell(2, ncols).value, c.value = c.value, None
    if last >= 2:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    t = ws["A1"]
    t.alignment = Alignment(horizontal="center", vertical="center")
    t.font = Font(bold=True, size=16, color="1F3864")
    for c in range(1, last + 1):
        ws.cell(1, c).fill = PatternFill("solid", fgColor="EAF1FB")
        ws.cell(1, c).border = Border(bottom=Side(style="medium", color="1F3864"))
    ws.row_dimensions[1].height = 32
    ws["A2"].font = Font(bold=True, size=10, color="404040")
    if ncols >= 2 and isinstance(ws.cell(2, ncols).value, str):
        ws.cell(2, ncols).alignment = Alignment(horizontal="right")
        ws.cell(2, ncols).font = Font(size=10, color="404040")
    if ws["A3"].value:
        ws["A3"].font = Font(size=10, italic=True, color="595959")


META_LABELS = {"모품목코드", "품목명", "BOM 버전", "기준수량", "사용여부", "작성 / 승인", "품목코드", "공정 버전", "작성"}


def meta_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(w, ws.column_dimensions[get_column_letter(i)].width or 0)


def meta_style(ws, rows=(3, 4)):
    """'모품목코드 | 값 | 품목명 | 값' 처럼 항목 이름과 값이 번갈아 있는 머리 정보 줄: 이름은 굵게 + 옅은 배경."""
    for r in rows:
        for c in ws[r]:
            if isinstance(c.value, str) and c.value in META_LABELS:
                c.font = Font(bold=True, size=10)
                c.fill = PatternFill("solid", fgColor="EDEDED")
                c.alignment = Alignment(horizontal="center")
            elif c.value is not None:
                c.alignment = Alignment(horizontal="left")


def header_row(ws, labels, fill="DDEBF7", widths=None, merge_to=None):
    ws.append(labels)
    r = ws.max_row
    for c in range(1, len(labels) + 1):
        cell = ws.cell(r, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[r].height = 24
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    # 틀 고정(freeze_panes)은 하지 않는다: 엑셀이 고정선을 표 밖 시트 끝까지 길게 그려서 테두리 밖으로 선이 삐져나온다
    ws._zebra_from = r
    ws._hdr_row = r
    ws._labels = labels
    title_block(ws, len(labels), merge_to)
    return r


def _text_w(v):
    """엑셀 열 너비 단위로 본 글자 길이: 한글·한자는 2, 그 밖은 1."""
    if isinstance(v, float):
        v = f"{v:g}"
    return sum(2.0 if ord(ch) >= 0x1100 else 1.0 for ch in str(v))


def _wrap_lines(text, width):
    """칸 너비(엑셀 단위)에서 낱말 단위로 줄을 바꿀 때 몇 줄이 되는지."""
    lines, cur = 1, 0.0
    for word in text.split(" "):
        w = _text_w(word)
        if cur and cur + 1 + w > width:
            lines, cur = lines + 1, 0.0
        cur = cur + (1 if cur else 0) + w
        while cur > width:                                     # 한 낱말이 칸보다 길면 글자 단위로 접힌다
            lines, cur = lines + 1, cur - width
    return lines


def autofit(ws, min_w=6, max_w=58):
    hdr, last = getattr(ws, "_hdr_row", None), getattr(ws, "_last_data_row", None)
    if hdr is None or last is None:
        return
    need = {}
    for row in ws.iter_rows(min_row=hdr, max_row=last):
        for c in row:
            if c.value is not None:
                need[c.column] = max(need.get(c.column, 0), _text_w(c.value) * (1.1 if c.row <= hdr + 1 and c.font and c.font.bold else 1.0))
    for r in (3, 4):                                         # BOM·라우팅 머리 정보 줄: 항목 이름, 그리고 합치지 않은 칸의 값
        if not any(c.value in META_LABELS for c in ws[r]):   # 항목 이름이 없는 줄(조회기간·작업일자 같은 한 줄 글)은 열 너비에 넣지 않는다
            continue
        for c in ws[r]:
            if c.value is None:
                continue
            if c.value in META_LABELS:
                need[c.column] = max(need.get(c.column, 0), _text_w(c.value) * 1.1)
            elif c.coordinate not in ws.merged_cells and hdr > 4:
                need[c.column] = max(need.get(c.column, 0), _text_w(c.value))
    for col, w in need.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, round(w + 3.5)))
    for row in ws.iter_rows(min_row=hdr + 1, max_row=last):      # 줄바꿈 칸이 있는 줄은 높이를 늘린다 (안 하면 글이 칸 위아래로 삐져나옴)
        lines = max((_wrap_lines(str(c.value), ws.column_dimensions[get_column_letter(c.column)].width - 3)
                     for c in row if c.value is not None and c.alignment is not None and c.alignment.wrap_text), default=1)
        if lines > 1:
            ws.row_dimensions[row[0].row].height = 15 * lines + 2


def print_setup(ws):
    """인쇄·PDF 모양: 가로 A4, 한 쪽 너비에 맞춤(세로는 쪽 수 제한 없음), 표 머리글 줄은 쪽마다 반복, 아래 가운데 쪽 번호."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.25)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "&P / &N"
    if getattr(ws, "_hdr_row", None):
        ws.print_title_rows = f"{ws._hdr_row}:{ws._hdr_row}"


def save_wb(wb, path):
    for ws in wb.worksheets:
        autofit(ws)
        print_setup(ws)
    wb.save(path)


_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}( \d{2}:\d{2})?$|^\d{2}:\d{2}$")
CENTER_LABELS = {"No", "NO", "Lv", "순서", "단위", "외주", "조달", "구분", "상태", "사용", "S/N", "LOT", "계획정지", "대분류", "유형", "단계", "제조구분", "근무조", "라인번호",
                 "4M", "검사", "검사결과", "작업팀", "사용중지", "위치", "동시처리", "요일", "투입공정", "결과", "확인", "검증 결과", "L/T", "MOQ", "사유코드", "공정코드", "품목구분", "휴일목록"}


def body(ws, values):
    """데이터 한 줄: 숫자는 오른쪽, 날짜·짧은 코드는 가운데, 글은 왼쪽. 한 줄 걸러 옅은 줄무늬."""
    ws.append([None if v == "" else v for v in values])
    r = ws.max_row
    ws._last_data_row = r
    zebra = (r - getattr(ws, "_zebra_from", r)) % 2 == 0
    labels = getattr(ws, "_labels", None)         # 열 제목에 따라 정렬을 정한다 (같은 열은 모두 같은 정렬)
    for c, v in enumerate(values, 1):
        cell = ws.cell(r, c)
        cell.border = BOX
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
            if labels and c <= len(labels) and labels[c - 1] == "시간":
                cell.number_format = "0.0"
        elif isinstance(v, str) and (_DATE_RE.match(v) or (labels and c <= len(labels) and labels[c - 1] in CENTER_LABELS)):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=isinstance(v, str) and len(v) > 40)
        if zebra:
            cell.fill = PatternFill("solid", fgColor="F7F9FC")


def write_csv(path, rows, cols):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


# ═════════════════════════ 5. 11종 만들기 ═════════════════════════
NAMES = {"SMT라인": ("E1021", "김도현"), "펌웨어기록실": ("E1305", "윤서준"), "광학조립실": ("E2011", "이수민"), "조립1라인": ("E2005", "박지훈"),
         "방폭시험실": ("E3015", "최민호"), "기밀시험기": ("E3009", "정하늘"), "검사라인": ("E4007", "오세린"), "에이징실": ("E3002", "한지우"), "포장라인": ("E5003", "강나래")}
WC = {  # 작업장: (설명, 위치, 동시처리, 시간당비용, [(설비, 설명)])
    "SMT라인": ("SMT 실장 라인 (마운터·리플로우·AOI)", "1층", 1, 85000, [("마운터-1", "칩 마운터 1호기"), ("리플로우-1", "리플로우 오븐 1호기"), ("AOI-1", "자동 광학 검사기")]),
    "펌웨어기록실": ("펌웨어·인증서 기록 스테이션", "1층", 2, 30000, [("기록지그-1", "FW 기록 지그 1호"), ("기록지그-2", "FW 기록 지그 2호")]),
    "광학조립실": ("렌즈·센서 조립, 초점 조정 (클린부스)", "2층", 2, 52000, [("초점조정기-1", "초점·화각 조정기 1호"), ("초점조정기-2", "초점·화각 조정기 2호")]),
    "조립1라인": ("제품 조립 1라인 (동시 2대)", "2층", 2, 42000, [("토크드라이버-1", "전동 토크 드라이버")]),
    "방폭시험실": ("방폭 내압 시험", "1층", 1, 60000, [("내압시험기-1", "수압 내압 시험기")]),
    "기밀시험기": ("방수 기밀(누설) 시험", "2층", 1, 38000, [("누설시험기-1", "에어 누설 시험기")]),
    "검사라인": ("기능 검사 라인", "2층", 2, 40000, [("검사지그-1", "기능 검사 지그 1호"), ("검사지그-2", "기능 검사 지그 2호")]),
    "에이징실": ("번인·에이징 랙 (동시 12대)", "2층", 12, 15000, [("에이징랙-1", "에이징 랙 1호")]),
    "포장라인": ("포장·라벨 라인", "1층", 1, 28000, [("중량선별기-1", "중량 선별기")]),
}
CODES = [("D01", "설비 고장", "설비", "N", "Y", "설비가 갑자기 멈춤. 수리 필요"), ("D02", "계획 외 정비", "설비", "N", "Y", "고장은 아니나 급히 정비하느라 정지"),
         ("D03", "정기 점검", "설비", "Y", "Y", "보전 계획에 따른 점검"), ("D04", "기계 이상", "설비", "N", "N", "D01과 중복되어 사용중지"),
         ("D05", "정전 / 에어 이상", "설비", "N", "Y", "전원·에어 공급 중단"), ("M01", "자재 대기", "자재", "N", "Y", "투입할 자재가 라인에 없음"),
         ("M02", "자재 불량", "자재", "N", "Y", "투입 자재 불량으로 교환"), ("M03", "자재 교체", "자재", "Y", "Y", "릴·박스 소진으로 교체"),
         ("S01", "모델 교체", "준비", "Y", "Y", "다른 모델로 전환"), ("S02", "지그 / 금형 교체", "준비", "Y", "Y", "지그·금형 교체"),
         ("S03", "준비시간 과다", "준비", "N", "Y", "셋업이 표준 준비시간을 넘김"), ("Q01", "품질 이상 대응", "품질", "N", "Y", "불량 발생으로 원인 확인·재작업"),
         ("Q02", "초품 검사 대기", "품질", "Y", "Y", "작업 시작 전 첫 제품 검사"), ("W01", "작업자 부재 / 교대", "작업자", "Y", "Y", "인원 부족 또는 교대 시간"),
         ("W02", "작업자 실수", "작업자", "N", "Y", "조작·조립 실수로 정지"), ("W03", "식사 / 휴식", "작업자", "Y", "Y", "정해진 휴게 시간"),
         ("P01", "전 공정 지연", "공정", "N", "Y", "앞 공정에서 물량이 오지 않음"), ("P02", "후 공정 적체", "공정", "N", "Y", "다음 공정이 막혀 내보낼 수 없음"),
         ("E01", "계획 정지", "기타", "Y", "Y", "생산 계획이 없어 정지"), ("Z99", "기타", "기타", "N", "Y", "위에 없는 사유 (비고에 적을 것)")]
DT_MEMO = {"D01": ["마운터 노즐 막힘", "지그 센서 고장", "컨베이어 정지"], "D02": ["벨트 장력 조정", "지그 핀 교체"], "D05": ["순간 정전", "에어 압력 저하"],
           "M01": ["자재창고 불출 지연", "렌즈 미입고", "개스킷 미입고"], "M02": ["릴 자재 불량 교환", "하우징 도장 불량 교환"], "M03": ["릴 소진 교체"],
           "S02": ["검사 지그 교체"], "Q01": ["불량 원인 확인"], "Q02": ["초품 검사 대기"], "W02": ["조립 순서 오류 수정"], "P01": ["앞 공정 물량 지연"], "P02": ["포장 적체"],
           "I01": ["인증서 서버 응답 없음"], "I02": ["새 FW 이미지 배포 대기"], "D06": ["누설 시험기 교정"], "D07": ["압력계 점검"], "D08": ["랙 온도 상한 초과"]}
DT_POOL = {"SMT라인": ["D01", "M02", "M03", "D02"], "펌웨어기록실": ["I01", "I02"], "광학조립실": ["M01", "Q01", "W02"], "조립1라인": ["M01", "W02", "P01"],
           "방폭시험실": ["D07", "Q02"], "기밀시험기": ["D06", "Q01"], "검사라인": ["Q02", "D01", "S02"], "에이징실": ["D08", "D05"], "포장라인": ["M01", "P02"]}
LOSS = {"OP-SMT": 0.010, "OP-FW": 0.003, "OP-OPT": 0.015, "OP-PT": 0.008, "OP-ASM": 0.008, "OP-EXP": 0.010, "OP-SEAL": 0.012, "OP-FCT": 0.015, "OP-AGE": 0.005, "OP-PKG": 0.002}
MAKE_OP = {"SMT": "OP-SMT", "광학 조립": "OP-OPT", "팬/틸트 조립": "OP-PT", "하우징 조립": "OP-ASM"}


def build(specs, out, seed=0, formats=OUT_FORMATS):
    rng = random.Random(seed or zlib.crc32("|".join(s["model"] for s in specs).encode()))
    out.mkdir(parents=True, exist_ok=True)
    products = [dict(spec=s, model=s["model"], **design(s, rng)) for s in specs]
    spec_by = {p["model"]: p["spec"] for p in products}

    # ── ① 품목 마스터 ──
    items, item_by = [], {}
    vendors = ["가상부품(주)", "테스트일렉", "샘플광학", "모의전자", "예시정밀"]

    def add_item(code, name, spec_text, kind, source, lot, serial, bom_, note, bs, vendor="", mpn="", use="Y"):
        if code in item_by:
            return
        made = date(2026, 3, 2) + timedelta(days=rng.randint(0, 180))
        it = {"품목코드": code, "품목명": name, "규격": spec_text, "품목구분": kind, "단위": "EA", "조달구분": source, "LOT관리": lot, "시리얼관리": serial,
              "안전재고": {"완제품": rng.choice([20, 30, 50]), "반제품": rng.choice([20, 50]), "부품": rng.choice([100, 200, 500, 1000])}[kind],
              "리드타임(일)": {"완제품": 7, "반제품": 3, "부품": rng.choice([7, 14, 21, 28])}[kind],
              "최소발주량": 1 if kind != "부품" else rng.choice([100, 500, 1000]), "주거래처(제조사)": vendor, "제조사품번": mpn, "기본BOM": bom_, "사용여부": use,
              "등록일": ds(made), "수정일": ds(min(date(2026, 10, 20), made + timedelta(days=rng.randint(0, 200)))), "비고": note, "basis": bs}
        items.append(it)
        item_by[code] = it

    for p in products:
        s, m = p["spec"], p["model"]
        wt, dims, col = sv(s, "weight_g", 0), sv(s, "dims"), sv(s, "color")
        note = " / ".join(x for x in [f"무게 {wt:g}g" if wt else "", f"치수 {dims}" if dims else "", f"색상 {col}" if col else "",
                                      "방수 " + ", ".join(s["ip_list"]) if s["ip_list"] else "", "충격 " + ", ".join(s["ik_list"]) if s["ik_list"] else ""] if x)
        if len(s["ik_list"]) > 1:
            note += f" (사양서에 충격 등급 {'·'.join(s['ik_list'])} 이 함께 적힘)"
        if s.get("maker"):
            note = " / ".join(x for x in [note, f"제조 및 A/S: {sv(s, 'maker')}"] if x)
        bs = " ".join(x for x in [f"{basis(s, 'resolution')} 해상도" if s["resolution"] else "", f"{basis(s, 'weight_g')} 무게" if wt else "",
                                  f"{basis(s, 'dims')} 치수" if dims else "", f"{basis(s, 'color')} 색상" if col else "",
                                  f"{basis(s, 'maker')} 제조사" if s.get("maker") else ""] if x) or "[사양서 추론]"
        add_item(m, s["title"] or m, " ".join(x for x in [sv(s, "resolution"), sv(s, "sensor")] if x), "완제품", "생산", "Y", "Y", f"BOM-{m}", note, bs)
        p["sf"] = {}
        for i, (key, name, proc, qty, bs2) in enumerate(p["semis"], 1):
            code = f"SF-{m}-{i:02d}"
            p["sf"][key] = code
            add_item(code, name, "", "반제품", "생산", "Y", "N", f"BOM-{code}", "", bs2)
        p["rm"] = []
        for key, name, spc, q, u, loss, proc, bs3 in p["parts"]:
            code = rm_code(name, spc)
            if code in item_by and (item_by[code]["품목명"], item_by[code]["규격"]) != (name, spc):
                raise RuntimeError(f"품목코드 충돌: {code} {name} {spc}")
            add_item(code, name, spc, "부품", "구매", "N", "N", "", rng.choice(["", "", "", "수입검사 필수", "정전기 주의"]) if proc == "SMT" else "", bs3,
                     vendor=vendors[zlib.crc32(name.encode()) % len(vendors)], mpn=f"MPN-{code[3:]}")
            p["rm"].append(code)
    for nm, spc in [("영상처리 SoC", "이전 세대"), ("플래시 메모리", "단종 용량"), ("체결 나사", "M3x6 SUS")]:   # BOM 에서 쓰지 않는 사용중지 품목
        code = rm_code(nm, spc)
        add_item(code, nm, spc, "부품", "구매", "N", "N", "", rng.choice(["단종(EOL)", "대체품 사용", "2025 사용중지"]), "[가상] 사용중지 품목 예시",
                 vendor=rng.choice(vendors), mpn=f"MPN-{code[3:]}", use="N")
    item_cols = ["품목코드", "품목명", "규격", "품목구분", "단위", "조달구분", "LOT관리", "시리얼관리", "안전재고", "리드타임(일)", "최소발주량",
                 "주거래처(제조사)", "제조사품번", "기본BOM", "사용여부", "등록일", "수정일", "비고"]
    write_csv(out / "품목마스터" / "품목마스터_정답.csv", items, item_cols + ["basis"])
    wb = Workbook()
    ws = wb.active
    ws.title = "품목리스트"
    head(ws, "품 목 등 록 현 황", "출력일: 2026-10-24  /  조회조건: 사용여부 전체", f"총 {len(items)}건", size=16)
    header_row(ws, ["No", "품목코드", "품목명", "규격", "구분", "단위", "조달", "LOT", "S/N", "안전재고", "L/T", "MOQ", "주거래처", "제조사품번", "기본BOM", "사용", "등록일", "수정일", "비고"],
               widths=[6, 22, 30, 34, 8, 6, 7, 6, 6, 9, 6, 8, 14, 14, 24, 6, 12, 12, 48])
    for i, it in enumerate(items, 1):
        body(ws, [i] + [it[c] for c in item_cols])
    save_wb(wb, out / "품목마스터" / "품목마스터_원본.xlsx")

    # ── ② BOM (정전개: 1단계 반제품 다음에 그 부품, 마지막에 완제품 직속 부품) ──
    bom = []
    wb = Workbook()
    wb.remove(wb.active)
    for p in products:
        m = p["model"]
        rows, no = [], 0
        for key, name, proc, qty, bs in p["semis"]:
            no += 1
            rows.append(dict(product=m, rev="Rev.A", no=no, level=1, parent=m, child=p["sf"][key], qty=qty, unit="EA", loss=0, process=proc, memo="", basis=bs))
            for (k2, nm, spc, q, u, loss, prc, bs2), code in zip(p["parts"], p["rm"]):
                if k2 == key:
                    no += 1
                    remark = "수입검사" if any(x in nm for x in ("센서", "검출기", "렌즈", "SoC", "TPM")) else ("대체품 가능" if any(x in nm for x in ("나사", "박스", "설명서")) else "")
                    rows.append(dict(product=m, rev="Rev.A", no=no, level=2, parent=p["sf"][key], child=code, qty=q, unit=u, loss=loss, process=prc, memo=remark,
                                     basis=bs2 + (" / 비고 [가상]" if remark else "")))
        for (k2, nm, spc, q, u, loss, prc, bs2), code in zip(p["parts"], p["rm"]):
            if k2 is None:
                no += 1
                rows.append(dict(product=m, rev="Rev.A", no=no, level=1, parent=m, child=code, qty=q, unit=u, loss=loss, process=prc, memo="", basis=bs2))
        bom += rows
        ws = wb.create_sheet(m[:31])
        ws["A1"] = "B O M 등 록 현 황"
        ws["A1"].font = Font(bold=True, size=15)
        ws["A2"], ws["H2"] = COMPANY, "출력일: 2026-10-25"
        ws.append(["모품목코드", m, None, "품목명", item_by[m]["품목명"], "BOM 버전", "Rev.A"])
        ws.append(["기준수량", 1, None, "사용여부", "사용", "작성 / 승인", "생산기술팀 / 생산관리팀"])
        for rng_ in ("B3:C3", "G3:H3", "B4:C4", "G4:K4"):
            ws.merge_cells(rng_)
        ws.append([])
        header_row(ws, ["No", "Lv", "품목코드", "품명", "규격", "소요량", "단위", "손실률(%)", "투입공정", "조달", "비고"], widths=[11, 6, 22, 28, 40, 12, 8, 10, 12, 7, 14])
        meta_style(ws)
        for r in rows:
            it = item_by[r["child"]]
            body(ws, [r["no"], str(r["level"]), r["child"], it["품목명"], it["규격"], r["qty"], r["unit"], r["loss"] or "", r["process"], it["조달구분"], r["memo"]])
    write_csv(out / "BOM" / "BOM_정답.csv", bom, ["product", "rev", "no", "level", "parent", "child", "qty", "unit", "loss", "process", "memo", "basis"])
    save_wb(wb, out / "BOM" / "BOM_원본.xlsx")

    # ── ③ 라우팅 ──
    route = []
    wb = Workbook()
    wb.remove(wb.active)
    for p in products:
        m = p["model"]
        ws = wb.create_sheet(m[:31])
        ws["A1"] = "공 정 표 (라우팅)"
        ws["A1"].font = Font(bold=True, size=15)
        ws["A2"], ws["H2"] = COMPANY, "출력일: 2026-10-26"
        ws.append(["품목코드", m, None, "품목명", item_by[m]["품목명"]])
        ws.append(["기준수량", 1, None, "작성", "생산기술팀", None, "공정 버전", "R1"])
        for rng_ in ("B3:C3", "E3:H3", "B4:C4", "E4:F4"):
            ws.merge_cells(rng_)
        ws.append([])
        header_row(ws, ["순서", "공정코드", "공정명", "작업장", "표준시간(분/EA)", "준비시간(분)", "투입인원", "공정검사항목", "외주", "비고"], widths=[11, 12, 22, 14, 14, 12, 12, 60, 6, 12])
        meta_style(ws)
        for i, o in enumerate(p["routing"], 1):
            r = dict(product=m, seq=i * 10, code=o["code"], name=o["name"], workcenter=o["wc"], std_min=o["std"], setup_min=o["setup"], crew=o["crew"],
                     inspection_item=o["insp"], outsource="N", note="", basis=o["basis"])
            route.append(r)
            body(ws, [r["seq"], r["code"], r["name"], r["workcenter"], r["std_min"], r["setup_min"], r["crew"], r["inspection_item"], "N", ""])
    write_csv(out / "라우팅" / "라우팅_정답.csv", route, ["product", "seq", "code", "name", "workcenter", "std_min", "setup_min", "crew", "inspection_item", "outsource", "note", "basis"])
    save_wb(wb, out / "라우팅" / "라우팅_원본.xlsx")

    # ── ④ 작업장 ──
    used = []
    for r in route:
        if r["workcenter"] not in used:
            used.append(r["workcenter"])
    std_by_wc = defaultdict(list)
    for r in route:
        std_by_wc[r["workcenter"]].append(r["std_min"])
    wsr, hours = [], []
    for wc in used:
        d_, loc, par_, cost, eq = WC[wc]
        wsr.append(dict(name=wc, desc=d_, wtype="라인", parent="", floor=loc, parallel=par_, capacity=round(par_ * 60 / (sum(std_by_wc[wc]) / len(std_by_wc[wc]))),
                        cost=cost, status="가동", holiday="2026 공장휴일", disabled="N", basis="[사양서 추론] 라우팅 공정의 작업장 / 동시처리·능력·비용 [가상]"))
        for e, ed in eq:
            wsr.append(dict(name=e, desc=ed, wtype="설비", parent=wc, floor=loc, parallel=1, capacity="", cost="", status="가동", holiday="2026 공장휴일", disabled="N", basis="[가상]"))
        for a, b in WORK:
            hours.append(dict(name=wc, start=f"{a // 60:02d}:00", end=f"{b // 60:02d}:00", enabled="Y", hours="4.0"))
    wsr.append(dict(name="조립2라인", desc="제품 조립 2라인 (증설 예정)", wtype="라인", parent="", floor="2층", parallel=2, capacity="", cost=42000, status="중지",
                    holiday="2026 공장휴일", disabled="Y", basis="[가상] 사용중지 작업장 예시"))
    write_csv(out / "작업장" / "작업장_정답.csv", wsr, ["name", "desc", "wtype", "parent", "floor", "parallel", "capacity", "cost", "status", "holiday", "disabled", "basis"])
    write_csv(out / "작업장" / "작업장_근무시간_정답.csv", hours, ["name", "start", "end", "enabled", "hours"])
    wb = Workbook()
    ws = wb.active
    ws.title = "작업장목록"
    ws["A1"] = "작 업 장 ( 라 인 ) 현 황"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"], ws["H2"] = COMPANY, "출력일: 2026-10-27"
    ws.append([])
    header_row(ws, ["No", "작업장명", "설명", "유형", "상위작업장", "위치", "동시처리", "능력(대/시)", "시간당비용(원)", "상태", "휴일목록", "사용중지"], widths=[6, 14, 36, 6, 12, 6, 9, 11, 13, 6, 14, 8])
    for i, r in enumerate(wsr, 1):
        body(ws, [i, r["name"], r["desc"], r["wtype"], r["parent"], r["floor"], r["parallel"], r["capacity"], r["cost"], r["status"], r["holiday"], r["disabled"]])
    w2 = wb.create_sheet("근무시간")
    head(w2, "작 업 장 근 무 시 간", "출력일: 2026-10-27", "평일(월~금) 기준, 점심 12:00~13:00 제외")
    header_row(w2, ["작업장명", "시작", "종료", "사용", "시간"], widths=[14, 8, 8, 6, 8])
    for r in hours:
        body(w2, [r["name"], r["start"], r["end"], r["enabled"], float(r["hours"])])
    save_wb(wb, out / "작업장" / "작업장_원본.xlsx")

    # ── ⑤ 비가동 사유 (기존 공통코드 20개 + 이 공장 공정에만 있는 코드) ──
    codes = list(CODES)
    if "펌웨어기록실" in used:
        codes += [("I01", "인증서 서버 연결 끊김", "정보", "N", "Y", "FW·인증서 기록 중 서버 연결 끊김"), ("I02", "FW 이미지 대기", "정보", "N", "Y", "새 펌웨어 이미지 배포 대기")]
    if "기밀시험기" in used:
        codes.append(("D06", "기밀 시험기 교정", "설비", "Y", "Y", "누설 시험기 정기 교정"))
    if "방폭시험실" in used:
        codes.append(("D07", "내압 시험기 점검", "설비", "Y", "Y", "방폭 내압 시험기 압력계 점검"))
    codes.append(("D08", "에이징실 온도 이탈", "설비", "N", "Y", "에이징 랙 온도가 설정 범위를 벗어남"))
    added = {"I01", "I02", "D06", "D07", "D08"}
    code_rows = [dict(code=c, name=n, group=g, planned=p_, use=u, description=d_,
                      basis="[가상] 이 공장 공정에서 추가한 코드 (공개 자료는 정지유형 등록 기능만 확인)" if c in added else "[가상] 기존 공통코드 표") for c, n, g, p_, u, d_ in codes]
    write_csv(out / "비가동사유" / "비가동사유_정답.csv", code_rows, ["code", "name", "group", "planned", "use", "description", "basis"])
    wb = Workbook()
    ws = wb.active
    ws.title = "비가동사유"
    head(ws, "공 통 코 드 : 비 가 동 사 유", "출력일: 2026-10-27", f"총 {len(code_rows)}건")
    header_row(ws, ["코드", "사유명", "대분류", "계획정지", "사용", "설명"], widths=[8, 22, 8, 9, 6, 40])
    for r in code_rows:
        body(ws, [r["code"], r["name"], r["group"], r["planned"], r["use"], r["description"]])
    save_wb(wb, out / "비가동사유" / "비가동사유_원본.xlsx")

    # ── ⑥ MPS (주별, 납기 금요일) ──
    mps = []
    fridays = [FIRST_MONDAY + timedelta(days=4 + 7 * w) for w in range(WEEKS)]
    for p in products:
        wkg = (sv(p["spec"], "weight_g", 0) or 500) / 1000
        base = 200 if wkg < 1 else 120 if wkg < 3 else 60 if wkg < 8 else 20
        weeks = sorted(rng.sample(range(WEEKS), k=max(4, WEEKS - 2 * (len(products) - 1)) if len(products) < 3 else 4))
        for w in weeks:
            qty = max(10, int(round(base * rng.uniform(0.5, 1.4) / 10)) * 10)
            lead = item_by[p["model"]]["리드타임(일)"]
            mps.append(dict(item=p["model"], name=item_by[p["model"]]["품목명"], bom=f"BOM-{p['model']}", qty=qty, unit="EA", warehouse="완제품창고",
                            delivery=ds(fridays[w]), release=ds(fridays[w] - timedelta(days=lead)), lead_days=lead, basis="[가상] 사양서에는 수량·납기가 없다"))
    mps.sort(key=lambda r: (r["delivery"], r["item"]))
    write_csv(out / "MPS" / "MPS_정답.csv", mps, ["item", "name", "bom", "qty", "unit", "warehouse", "delivery", "release", "lead_days", "basis"])
    wb = Workbook()
    ws = wb.active
    ws.title = "MPS목록"
    head(ws, "기 준 생 산 계 획 (MPS)", "출력일: 2026-10-28", f"계획기간: {ds(FIRST_MONDAY)} ~ {ds(FIRST_MONDAY + timedelta(days=7 * WEEKS - 1))}")
    header_row(ws, ["No", "품목코드", "품명", "BOM", "계획수량", "단위", "입고창고", "납기일", "발주시작일", "L/T(일)"], widths=[6, 16, 30, 22, 9, 6, 11, 12, 12, 8])
    for i, r in enumerate(mps, 1):
        body(ws, [i, r["item"], r["name"], r["bom"], r["qty"], r["unit"], r["warehouse"], r["delivery"], r["release"], r["lead_days"]])
    w2 = wb.create_sheet("주간보기")
    w2["A1"] = "주 간 생 산 계 획"
    w2["A1"].font = Font(bold=True, size=15)
    w2["A2"] = f"{COMPANY}    단위: EA"
    w2.append([])
    mondays = [FIRST_MONDAY + timedelta(weeks=w) for w in range(WEEKS)]
    header_row(w2, ["품목코드", "품명"] + [f"{d.month}/{d.day}주" for d in mondays] + ["합계"], fill="E2EFDA", widths=[16, 30] + [8] * WEEKS + [8])
    for p in products:
        vals = [sum(r["qty"] for r in mps if r["item"] == p["model"] and r["delivery"] == ds(fridays[w])) or "" for w in range(WEEKS)]
        body(w2, [p["model"], item_by[p["model"]]["품목명"]] + vals + [sum(v for v in vals if v)])
    save_wb(wb, out / "MPS" / "MPS_원본.xlsx")

    # ── ⑦ 생산계획 (MPS 1줄 = 계획 1줄, 같은 주는 같은 계획번호) ──
    plans, week_no = [], {}
    for r in mps:
        fri = date.fromisoformat(r["delivery"])
        week_no.setdefault(r["delivery"], len(week_no) + 1)
        mon = fri - timedelta(days=4)
        plans.append(dict(plan_no=f"PP-{mon.strftime('%y%m')}-{week_no[r['delivery']]:02d}", plan_date=ds(mon - timedelta(days=7)), item=r["item"], name=r["name"],
                          bom=r["bom"], qty=r["qty"], unit="EA", start=ds(mon), end=r["delivery"], warehouse="완제품창고", remain="", produced="",
                          basis="[가상] MPS 를 주 단위 계획으로 / 잔량·생산수량은 포장 실적"))

    # ── ⑧ 작업지시 + 일정 (작업장마다 동시처리 수만큼 자리를 두고 가장 먼저 비는 자리에 넣는다) ──
    route_by = defaultdict(list)
    for r in route:
        route_by[r["product"]].append(r)
    par = {r["name"]: int(r["parallel"]) for r in wsr if r["wtype"] == "라인"}
    slots = {wc: [[datetime.combine(FIRST_MONDAY, time(8)), None] for _ in range(par[wc])] for wc in used}
    wos, wo_ops, changeovers = [], [], []
    n = 0
    for pl in plans:
        left = pl["qty"]
        while left > 0:
            q = min(LOT_MAX, left)
            left -= q
            n += 1
            wo = f"WO-{date.fromisoformat(pl['start']).strftime('%y%m')}-{n:03d}"
            prev_end = datetime.combine(date.fromisoformat(pl["start"]), time(8))
            qty_in, rows = q, []
            for r in route_by[pl["item"]]:
                wc = r["workcenter"]
                slot = min(slots[wc], key=lambda x: x[0])
                start = next_work(max(prev_end, slot[0]))
                if slot[1] and slot[1] != pl["item"]:
                    changeovers.append(dict(t=start, wc=wc, a=slot[1], b=pl["item"], wo=wo, op=r["name"]))
                end = add_work(start, math.ceil(r["setup_min"] + q * r["std_min"]))
                slot[0], slot[1] = end, pl["item"]
                prev_end = end
                if start < TODAY:
                    ratio = 1.0 if end <= TODAY else wmin(start, TODAY) / max(1, wmin(start, end))
                    processed = min(qty_in, round(qty_in * ratio))
                    loss = min(processed, round(processed * LOSS[r["code"]] * rng.uniform(0.4, 1.6)))
                    status = "완료" if end <= TODAY else "진행"
                else:
                    processed, loss, status = 0, 0, "대기"
                rows.append(dict(wo=wo, seq=r["seq"], name=r["name"], wc=wc, start=dts(start), end=dts(end), std=r["std_min"], done=processed - loss, loss=loss,
                                 status=status, code=r["code"], qty_in=qty_in, basis="[가상] 일정·수량 / 순서·작업장·표준시간은 라우팅"))
                if status == "완료":
                    qty_in = processed - loss
            wo_ops += rows
            st = "완료" if all(o["status"] == "완료" for o in rows) else "미시작" if all(o["status"] == "대기" for o in rows) else "진행중"
            first = datetime.fromisoformat(rows[0]["start"])
            wos.append(dict(wo=wo, issue=ds(first.date() - timedelta(days=3)), plan_no=pl["plan_no"], item=pl["item"], name=pl["name"], bom=pl["bom"], qty=q, unit="EA",
                            status=st, start=ds(first.date()), end=rows[-1]["end"][:10], wip="재공창고", fg="완제품창고", lot=f"L{wo[3:]}",
                            produced=rows[-1]["done"] if st == "완료" else "", note=rng.choice(["", "", "", "", "고객 납기 우선"]),
                            basis="[가상] 생산계획을 LOT(최대 200) 로 나눔"))
    wo_by = {w["wo"]: w for w in wos}
    for pl in plans:
        prod = sum(o["done"] for o in wo_ops if o["code"] == "OP-PKG" and wo_by[o["wo"]]["plan_no"] == pl["plan_no"] and wo_by[o["wo"]]["item"] == pl["item"])
        pl["produced"], pl["remain"] = prod, pl["qty"] - prod
    subs = []
    for pl in plans:
        p = next(x for x in products if x["model"] == pl["item"])
        for key, name, proc, qty, bs in p["semis"]:
            rows = [o for o in wo_ops if o["code"] == MAKE_OP[proc] and wo_by[o["wo"]]["plan_no"] == pl["plan_no"] and wo_by[o["wo"]]["item"] == pl["item"]]
            subs.append(dict(plan_no=pl["plan_no"], sub=p["sf"][key], sub_name=name, parent=pl["item"], bom=f"BOM-{p['sf'][key]}", level=1, need=qty * pl["qty"], unit="EA",
                             mtype="자체", vendor="", start=min(o["start"] for o in rows)[:10], end=max(o["end"] for o in rows)[:10],
                             produced=sum(o["done"] for o in rows) * qty, basis="[사양서 추론] BOM 1단계 × 계획수량 / 기간·생산수량은 작업지시 공정"))
    write_csv(out / "생산계획" / "생산계획_정답.csv", plans, ["plan_no", "plan_date", "item", "name", "bom", "qty", "unit", "start", "end", "warehouse", "remain", "produced", "basis"])
    write_csv(out / "생산계획" / "반제품소요_정답.csv", subs, ["plan_no", "sub", "sub_name", "parent", "bom", "level", "need", "unit", "mtype", "vendor", "start", "end", "produced", "basis"])
    wb = Workbook()
    ws = wb.active
    ws.title = "생산계획"
    head(ws, "생 산 계 획 현 황", "출력일: 2026-10-29", f"조회: {ds(FIRST_MONDAY)} ~ {ds(FIRST_MONDAY + timedelta(days=7 * WEEKS - 1))}")
    header_row(ws, ["계획번호", "계획일", "품목코드", "품명", "BOM", "계획수량", "단위", "계획시작", "계획종료", "입고창고", "잔량", "생산수량"], widths=[12, 12, 16, 30, 22, 9, 6, 12, 12, 11, 7, 9])
    for r in plans:
        body(ws, [r["plan_no"], r["plan_date"], r["item"], r["name"], r["bom"], r["qty"], r["unit"], r["start"], r["end"], r["warehouse"], r["remain"], r["produced"]])
    w2 = wb.create_sheet("반제품소요")
    head(w2, "반 제 품 소 요 계 획", "출력일: 2026-10-29", "생산계획 번호별, BOM 1단계 모듈 기준")
    header_row(w2, ["계획번호", "반제품코드", "반제품명", "완제품", "BOM", "단계", "필요수량", "단위", "제조구분", "공급사", "시작일", "종료일", "생산수량"], widths=[12, 22, 16, 16, 26, 6, 9, 6, 9, 8, 12, 12, 9])
    for r in subs:
        body(w2, [r["plan_no"], r["sub"], r["sub_name"], r["parent"], r["bom"], r["level"], r["need"], r["unit"], r["mtype"], r["vendor"], r["start"], r["end"], r["produced"]])
    save_wb(wb, out / "생산계획" / "생산계획_원본.xlsx")

    wo_items = []
    for w in wos:
        for b in bom:
            if b["product"] == w["item"] and b["level"] == 1:
                op = next(o for o in wo_ops if o["wo"] == w["wo"] and o["name"] == b["process"])
                need = b["qty"] * w["qty"]
                wo_items.append(dict(wo=w["wo"], process=b["process"], code=b["child"], name=item_by[b["child"]]["품목명"], need=need,
                                     issued=need if op["status"] != "대기" else 0, warehouse="재공창고", used=(op["done"] + op["loss"]) * b["qty"],
                                     basis="[사양서 추론] BOM 1단계 × 지시수량 / 불출·사용은 공정 실적"))
    write_csv(out / "작업지시서" / "작업지시_정답.csv", wos, ["wo", "issue", "plan_no", "item", "name", "bom", "qty", "unit", "status", "start", "end", "wip", "fg", "lot", "produced", "note", "basis"])
    write_csv(out / "작업지시서" / "작업지시_자재_정답.csv", wo_items, ["wo", "process", "code", "name", "need", "issued", "warehouse", "used", "basis"])
    write_csv(out / "작업지시서" / "작업지시_공정_정답.csv", wo_ops, ["wo", "seq", "name", "wc", "start", "end", "std", "done", "loss", "status", "code", "qty_in", "basis"])
    wb = Workbook()
    ws = wb.active
    ws.title = "작업지시"
    head(ws, "작 업 지 시 현 황", "출력일: 2026-11-12", f"지시 {len(wos)}건")
    header_row(ws, ["작업지시번호", "지시일", "생산계획", "품목코드", "품명", "BOM", "지시수량", "단위", "상태", "계획시작", "계획종료", "재공창고", "완제품창고", "LOT", "생산수량", "비고"],
               widths=[13, 11, 12, 16, 28, 22, 9, 6, 8, 11, 11, 10, 10, 11, 9, 14])
    for r in wos:
        body(ws, [r["wo"], r["issue"], r["plan_no"], r["item"], r["name"], r["bom"], r["qty"], r["unit"], r["status"], r["start"], r["end"], r["wip"], r["fg"], r["lot"], r["produced"], r["note"]])
    w2 = wb.create_sheet("소요자재")
    head(w2, "작 업 지 시 별 소 요 자 재", "출력일: 2026-11-12", "BOM 1단계 기준")
    header_row(w2, ["작업지시번호", "투입공정", "자재코드", "자재명", "필요수량", "출고수량", "출고창고", "소모수량"], widths=[13, 12, 22, 18, 9, 9, 10, 9])
    for r in wo_items:
        body(w2, [r["wo"], r["process"], r["code"], r["name"], r["need"], r["issued"], r["warehouse"], r["used"]])
    w3 = wb.create_sheet("공정")
    head(w3, "작 업 지 시 별 공 정", "출력일: 2026-11-12", "라우팅 순서 기준")
    header_row(w3, ["작업지시번호", "순서", "공정", "작업장", "계획시작", "계획종료", "표준시간(분)", "완료수량", "손실수량", "상태"], widths=[13, 6, 14, 12, 16, 16, 12, 9, 9, 7])
    for r in wo_ops:
        body(w3, [r["wo"], r["seq"], r["name"], r["wc"], r["start"], r["end"], r["std"], r["done"], r["loss"], r["status"]])
    save_wb(wb, out / "작업지시서" / "작업지시서_원본.xlsx")

    # ── ⑨ 작업실적 (생산 작업 일지, 근무일마다 시트 1장) ──
    rt_by = {(r["product"], r["code"]): r for r in route}
    lines_ = []
    for o in wo_ops:
        if o["status"] == "대기":
            continue
        w = wo_by[o["wo"]]
        s0, e0 = datetime.fromisoformat(o["start"]), min(datetime.fromisoformat(o["end"]), TODAY)
        by_day = defaultdict(list)
        for a, b in segments(s0, e0):
            by_day[a.date()].append((a, b))
        days = sorted(by_day)
        total = sum(wmin(a, b) for d_ in days for a, b in by_day[d_]) or 1
        processed, lossall = o["done"] + o["loss"], o["loss"]
        left_p, left_l = processed, lossall
        for i, d_ in enumerate(days):
            mins = sum(round((b - a).total_seconds() / 60) for a, b in by_day[d_])
            last = i == len(days) - 1
            qp = left_p if last else min(left_p, round(processed * mins / total))
            ql = left_l if last else min(left_l, qp, round(lossall * mins / total))
            if qp <= 0:
                continue
            left_p -= qp
            left_l -= ql
            lines_.append(dict(o=o, w=w, day=d_, start=by_day[d_][0][0], end=by_day[d_][-1][1], mins=mins, qty=qp, loss=ql))
        if left_p or left_l:     # 마지막 날 수량이 0 이라 건너뛰었으면 앞 줄에 붙인다
            for x in reversed(lines_):
                if x["o"] is o:
                    x["qty"] += left_p
                    x["loss"] += left_l
                    break
    lines_.sort(key=lambda x: (x["day"], x["start"], x["w"]["wo"], x["o"]["seq"]))
    by_date = defaultdict(list)
    for x in lines_:
        by_date[x["day"]].append(x)
    log, serial_n = [], defaultdict(int)
    wb = Workbook()
    wb.remove(wb.active)
    total_plan, cum = sum(w["qty"] for w in wos), 0
    top = ["NO", "작업시간", None, "작업지시", "LOT", "품목코드", "품명", "작업장", "수량(EA)", "양품", "손실", "공정명", "작업팀", "작업자", "작업공수(분)", "설비가동(분)",
           "수율(%)", "품질공정검사", None, "시작 시리얼", "끝 시리얼", "FW 버전", "소비전력(W)", "중량(g)", "조치/비고"]
    for d_ in sorted(by_date):
        ws = wb.create_sheet(d_.strftime("%m%d"))
        ws["A1"] = "생 산 작 업 일 지"
        ws["A1"].font = Font(bold=True, size=16)
        ws["I1"], ws["J1"], ws["K1"] = "결재", "담당", "팀장"
        ws["A3"] = f"작업일자 : {d_.year}년 {d_.month}월 {d_.day}일 ({WEEKDAY[d_.weekday()]})"
        ws.append([])
        ws.append(top)
        ws.append([None, "시작", "종료"] + [None] * 14 + ["확인", "결과"] + [None] * 6)
        r1 = ws.max_row - 1
        ws.merge_cells(start_row=r1, start_column=2, end_row=r1, end_column=3)
        ws.merge_cells(start_row=r1, start_column=18, end_row=r1, end_column=19)
        for rr in (r1, r1 + 1):
            for c in range(1, len(top) + 1):
                cell = ws.cell(rr, c)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                cell.border = BOX
        for c in ("I1", "J1", "K1"):                           # 결재란
            ws[c].font = Font(bold=True, size=10)
            ws[c].alignment = Alignment(horizontal="center")
            ws[c].border = BOX
        title_block(ws, len(top), merge_to=8)
        ws["A3"].font = Font(bold=True, size=11)
        ws._zebra_from = r1 + 1
        ws._hdr_row = r1
        ws._labels = top
        produced = 0
        for no, x in enumerate(by_date[d_], 1):
            o, w = x["o"], x["w"]
            rt, sp = rt_by[(w["item"], o["code"])], spec_by[w["item"]]
            oid, onm = NAMES[o["wc"]]
            good = x["qty"] - x["loss"]
            team = "A조" if x["start"].hour < 12 else rng.choice(["A조", "B조"])
            check = "○" if o["code"] in ("OP-OPT", "OP-EXP", "OP-SEAL", "OP-FCT", "OP-PKG") else ""
            result = ("NG" if x["loss"] else "OK") if check else ""
            note = f"불량 {x['loss']}EA 폐기" if x["loss"] else rng.choice(["", "", "", "", "자재 교체"])
            sf = st_ = fw = power = weight = ""
            bs = "[가상] 수량·시간·작업자"
            if o["code"] == "OP-FCT" and good:
                sf = f"{w['item']}-{d_.strftime('%y%m')}{serial_n[w['item']] + 1:05d}"
                serial_n[w["item"]] += good
                st_ = f"{w['item']}-{d_.strftime('%y%m')}{serial_n[w['item']]:05d}"
                if fl(sp, "ethernet") or fl(sp, "tpm"):
                    fw = "1.00.00"
                bs = "[가상] 수량·시리얼·FW"
                if sp["power"]:
                    pw = sp["power"][0]
                    bad = x["loss"] > 0 and rng.random() < 0.4
                    power = round(pw["v"] * (1.04 if bad else rng.uniform(0.70, 0.93)), 2)
                    if bad:
                        note = f"소비전력 {power}W > 기준 {pw['v']:g}W, 불량 {x['loss']}EA 폐기"
                    bs = f"[사양서 p.{pw['p']}] 소비전력 기준 {pw['src']} {pw['v']:g}W / 측정값·시리얼 [가상]"
            if o["code"] == "OP-PKG" and good and sv(sp, "weight_g", 0):
                wt = sv(sp, "weight_g", 0)
                bad = x["loss"] > 0 and rng.random() < 0.5
                weight = round(wt * (1.05 if bad else rng.uniform(0.985, 1.015)), 1)
                if bad:
                    note = f"중량 {weight}g 기준 이탈(±3%), 불량 {x['loss']}EA 폐기"
                bs = f"{basis(sp, 'weight_g')} 제품 무게 {wt:g}g ±3% / 측정값 [가상]"
            labor = x["mins"] * rt["crew"]
            yld = round(100 * good / x["qty"], 1)
            body(ws, [no, x["start"].strftime("%H:%M"), x["end"].strftime("%H:%M"), w["wo"], w["lot"], w["item"], w["name"], o["wc"], x["qty"], good, x["loss"], o["name"],
                      team, f"{onm}({oid})", labor, x["mins"], yld, check, result, sf, st_, fw, power, weight, note])
            if o["code"] == "OP-PKG":
                produced += good
            log.append(dict(date=ds(d_), no=no, start=x["start"].strftime("%H:%M"), end=x["end"].strftime("%H:%M"), wo=w["wo"], lot=w["lot"], item=w["item"], name=w["name"],
                            qty=x["qty"], process=o["name"], workcenter=o["wc"], team=team, operator_id=oid, operator=onm, good_qty=good, loss=x["loss"], labor_minutes=labor,
                            machine_run_minutes=x["mins"], yield_pct=yld, qc_check=check, qc_result=result, serial_from=sf, serial_to=st_, fw_version=fw, power_w=power,
                            weight_g=weight, note=note, basis=bs))
        cum += produced
        ws.append([])
        ws.append(["금일생산량", None, produced, None, "누 계", cum, None, "출고량", None, None, "총잔량", None, total_plan - cum])
        for c in ws[ws.max_row]:
            if c.value is None:
                continue
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="right" if isinstance(c.value, int) else "center")
            if isinstance(c.value, str):
                c.fill = PatternFill("solid", fgColor="EDEDED")
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)   # A열(NO)은 좁아서 '금일생산량' 이 잘린다: A·B 칸을 합친다
        for i, wd in enumerate([5, 7, 7, 13, 11, 16, 26, 11, 8, 6, 6, 12, 7, 15, 11, 11, 8, 6, 6, 22, 22, 9, 10, 9, 30], 1):
            ws.column_dimensions[get_column_letter(i)].width = wd
    write_csv(out / "작업실적" / "작업실적_정답.csv", log, ["date", "no", "start", "end", "wo", "lot", "item", "name", "qty", "process", "workcenter", "team", "operator_id", "operator",
                                                       "good_qty", "loss", "labor_minutes", "machine_run_minutes", "yield_pct", "qc_check", "qc_result", "serial_from", "serial_to",
                                                       "fw_version", "power_w", "weight_g", "note", "basis"])
    save_wb(wb, out / "작업실적" / "작업실적_원본.xlsx")

    # ── ⑩ 비가동 기록: 실제로 일한 시간 안에서, 같은 작업장끼리 겹치지 않게 ──
    valid = {x[0] for x in codes if x[4] == "Y"}
    events, busy = [], defaultdict(list)
    for c in changeovers:
        st, en = c["t"], c["t"] + timedelta(minutes=20)
        if en <= TODAY and segments(st, en) == [(st, en)] and not any(not (en <= a or st >= b) for a, b in busy[c["wc"]]):
            events.append(dict(st=st, en=en, wc=c["wc"], code="S01", wo=c["wo"], memo=f"모델 교체 {c['a']} → {c['b']}"))
            busy[c["wc"]].append((st, en))
    for o in wo_ops:
        if o["status"] == "대기":
            continue
        s0, e0 = datetime.fromisoformat(o["start"]), min(datetime.fromisoformat(o["end"]), TODAY)
        if wmin(s0, e0) < 120 or rng.random() > 0.4:
            continue
        segs = [(a, b) for a, b in segments(s0, e0) if (b - a).total_seconds() >= 40 * 60]
        if not segs:
            continue
        a, b = rng.choice(segs)
        span = int((b - a).total_seconds() // 60)
        length = rng.choice([x for x in (10, 15, 20, 30, 45) if x <= span - 5])
        st = a + timedelta(minutes=rng.randint(0, span - length))
        en = st + timedelta(minutes=length)
        if any(not (en <= x0 or st >= x1) for x0, x1 in busy[o["wc"]]):
            continue
        c = rng.choice([c for c in DT_POOL[o["wc"]] if c in valid] or ["Z99"])
        busy[o["wc"]].append((st, en))
        events.append(dict(st=st, en=en, wc=o["wc"], code=c, wo=o["wo"], memo=rng.choice(DT_MEMO.get(c, ["확인 중"]))))
    events.sort(key=lambda e: (e["st"], e["wc"]))
    cname = {x[0]: x[1] for x in codes}
    dt_rows = [dict(date=ds(e["st"].date()), wc=e["wc"], start=dts(e["st"]), end=dts(e["en"]), minutes=int((e["en"] - e["st"]).total_seconds() // 60), code=e["code"],
                    reason=cname[e["code"]], operator_id=NAMES[e["wc"]][0], operator=NAMES[e["wc"]][1], wo=e["wo"], memo=e["memo"], basis="[가상] 발생 시각·사유") for e in events]
    write_csv(out / "비가동기록" / "비가동기록_정답.csv", dt_rows, ["date", "wc", "start", "end", "minutes", "code", "reason", "operator_id", "operator", "wo", "memo", "basis"])
    wb = Workbook()
    ws = wb.active
    ws.title = "비가동기록"
    head(ws, "설 비 비 가 동 기 록", "출력일: 2026-11-12", f"조회기간: {ds(FIRST_MONDAY)} ~ {ds(TODAY.date())}    건수: {len(dt_rows)}건    합계: {sum(r['minutes'] for r in dt_rows)}분")
    header_row(ws, ["No", "일자", "작업장", "시작", "종료", "비가동(분)", "사유코드", "사유", "작업자", "작업지시", "비고"], widths=[5, 11, 12, 7, 7, 10, 9, 20, 15, 13, 30])
    for i, r in enumerate(dt_rows, 1):
        body(ws, [i, r["date"], r["wc"], r["start"][11:], r["end"][11:], r["minutes"], r["code"], r["reason"], f"{r['operator']}({r['operator_id']})", r["wo"], r["memo"]])
    save_wb(wb, out / "비가동기록" / "비가동기록_원본.xlsx")

    # ── ⑪ 4M 변경점: 출력 기준일까지 근무일마다, 라인마다 1줄 이상 (변경 없으면 '없음') ──
    pool4 = []
    for p in products:
        sp = p["spec"]
        if not fl(sp, "mount") and (sp["sensor"] or sp["thermal"]):
            pool4.append(("광학조립실", "자재", "광학 조립 / 렌즈 공급 로트 변경", "초품 화각·초점 확인", "화각·초점", "초품 검사"))
        if fl(sp, "ir"):
            pool4.append(("SMT라인", "자재", "SMT / IR LED 자재 로트 변경", "점등·광량 확인", "IR 점등", "초품 검사"))
        if sp["ip_list"]:
            pool4.append(("기밀시험기", "설비", "기밀 시험 / 누설 시험기 교정 후 재투입", "기준 시료 확인", "누설량", "기준 시료 측정"))
            pool4.append(("조립1라인", "자재", "하우징 조립 / 방수 개스킷 로트 변경", "기밀 시험 전수 확인", "기밀", "전수 검사"))
        if fl(sp, "ethernet") or fl(sp, "tpm"):
            pool4.append(("펌웨어기록실", "방법", "FW 기록 / FW 버전 업데이트", "작업표준서 개정, 초품 검사", "FW 버전", "초품 검사"))
        if fl(sp, "motor_pt"):
            pool4.append(("조립1라인", "방법", "팬/틸트 조립 / 그리스 도포량 변경", "작업표준서 개정", "구동 소음", "초품 검사"))
    pool4 += [("광학조립실", "사람", "광학 조립 / 신규 작업자 투입", "교육 후 초품 확인", "작업 숙련도", "초품 검사"),
              ("SMT라인", "설비", "SMT / 마운터 노즐 교체", "초품 실장 상태 확인", "실장 상태", "AOI 확인"),
              ("검사라인", "설비", "기능 검사 / 검사 지그 교체", "기준 시료 확인", "지그 상태", "기준 시료 측정")]
    pool4 = list(dict.fromkeys(x for x in pool4 if x[0] in used))
    rng.shuffle(pool4)
    m4_rows = []
    wb = Workbook()
    wb.remove(wb.active)
    for d_ in workdays(FIRST_MONDAY, TODAY.date()):
        ws = wb.create_sheet(d_.strftime("%m%d"))
        ws["A1"] = "4M 변 경 점 기 록 표"
        ws["A1"].font = Font(bold=True, size=15)
        ws["A2"] = f"작성일 : {d_.year}.{d_.month:02d}.{d_.day:02d} ({WEEKDAY[d_.weekday()]})"
        ws["J2"] = "관리자 확인 :"
        ws.append([])
        header_row(ws, ["라인명", "라인번호", "부서·조", "근무조", "작성자(조장)", "4M", "변경 내용 (공정명 / 내용)", "조치·대책", "검증 항목", "검증 방법", "검증 결과", "검증 담당(QA)", "관리자 의견"],
                   fill="E2EFDA", widths=[12, 8, 16, 7, 11, 6, 42, 26, 12, 13, 8, 13, 14])
        for no, wc in enumerate(used, 1):
            rows = [("방법", f"{c['op']} / 모델 변경 {c['a']} → {c['b']}", "프로그램·작업표준서 교체, 셋업 확인", "셋업 상태", "초품 검사", "OK", "")
                    for c in changeovers if c["wc"] == wc and c["t"].date() == d_ and c["t"] < TODAY]
            cand = [x for x in pool4 if x[0] == wc]
            if cand and rng.random() < 0.2:
                pool4.remove(cand[0])
                _, m4, ch, act, vi, vm = cand[0]
                rows.append((m4, ch, act, vi, vm, rng.choice(["OK", "OK", "OK", "재확인"]), rng.choice(["", "", "다음 주 재점검"])))
            if not rows:
                rows.append(("", "없음", "", "", "", "", ""))
            for m4, ch, act, vi, vm, vr, mn in rows:
                body(ws, [wc, no, f"생산1팀 {wc}", "주간", NAMES[wc][1], m4, ch, act, vi, vm, vr, "송미라 (서명)" if vr else "", mn])
                m4_rows.append(dict(date=ds(d_), line=wc, line_no=no, team=f"생산1팀 {wc}", shift="주간", leader=NAMES[wc][1], m4=m4, change=ch, action=act, verify_item=vi,
                                    verify_method=vm, verify_result=vr, qa="송미라" if vr else "", manager_note=mn,
                                    basis="[가상] 변경 내용 / 소재는 [사양서 추론] 부품·공정" if m4 else "[가상] 변경 없음"))
    write_csv(out / "4M변경점" / "4M변경점_정답.csv", m4_rows, ["date", "line", "line_no", "team", "shift", "leader", "m4", "change", "action", "verify_item", "verify_method",
                                                          "verify_result", "qa", "manager_note", "basis"])
    save_wb(wb, out / "4M변경점" / "4M변경점_원본.xlsx")

    summary = [{k: v for k, v in p["spec"].items() if k != "flags"} | {"flags": {k: (v["p"] if v else None) for k, v in p["spec"]["flags"].items()}} for p in products]
    (out / "사양추출.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    if formats:
        export([out / f / f"{f}_원본.xlsx" for f in EXPORT_DOCS], formats)
    return out


# ═════════════════════════ 6. 검증 ═════════════════════════
def _rd(p):
    return list(csv.DictReader(open(p, encoding="utf-8-sig")))


def _xlsx_rows(path, first_label, sheets=None):
    """원본 xlsx 에서 머리글(첫 칸 = first_label) 아래 데이터 줄 수를 센다 (시트 여러 장이면 합)."""
    wb = openpyxl.load_workbook(path, read_only=True)
    total, gap = 0, []
    for name in wb.sheetnames:
        if sheets and name not in sheets:
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows) if r and r[0] == first_label), None)
        if h is None:
            continue
        if h + 1 < len(rows) and all(v is None for v in rows[h + 1]) and any(any(v is not None for v in r) for r in rows[h + 2:]):
            gap.append(name)                       # 표 머리글 바로 아래가 빈 줄 (데이터가 있는데도)
        started = False
        for r in rows[h + 1:]:
            a = r[0] if r else None
            if a in (None, ""):
                if started:
                    break
                continue
            if a in ("금일생산량", "합계"):
                break
            started = True
            total += 1
    wb.close()
    return total, gap


def verify(out, formats=OUT_FORMATS):
    E = []
    for f in FOLDERS:
        if not (out / f / f"{f}_원본.xlsx").exists():
            E.append(f"{f}: 원본.xlsx 없음")
    need = ["품목마스터/품목마스터_정답.csv", "BOM/BOM_정답.csv", "라우팅/라우팅_정답.csv", "작업장/작업장_정답.csv", "작업장/작업장_근무시간_정답.csv", "비가동사유/비가동사유_정답.csv",
            "MPS/MPS_정답.csv", "생산계획/생산계획_정답.csv", "생산계획/반제품소요_정답.csv", "작업지시서/작업지시_정답.csv", "작업지시서/작업지시_자재_정답.csv",
            "작업지시서/작업지시_공정_정답.csv", "작업실적/작업실적_정답.csv", "비가동기록/비가동기록_정답.csv", "4M변경점/4M변경점_정답.csv"]
    for n_ in need:
        if not (out / n_).exists():
            E.append(f"{n_} 없음")
    if E:
        return E, []
    C = {n_.split("/")[1][:-4]: _rd(out / n_) for n_ in need}
    items = {r["품목코드"]: r for r in C["품목마스터_정답"]}
    bom, route = C["BOM_정답"], C["라우팅_정답"]
    ws = {r["name"]: r for r in C["작업장_정답"]}
    codes = {r["code"]: r for r in C["비가동사유_정답"]}
    mps, plans, subs = C["MPS_정답"], C["생산계획_정답"], C["반제품소요_정답"]
    wos = {r["wo"]: r for r in C["작업지시_정답"]}
    wit, wop, log, dt, m4 = C["작업지시_자재_정답"], C["작업지시_공정_정답"], C["작업실적_정답"], C["비가동기록_정답"], C["4M변경점_정답"]
    T = TODAY

    # (1) 원본 xlsx 와 정답 csv 의 줄 수가 같다
    for path, label, sheets, rows in [("품목마스터/품목마스터_원본.xlsx", "No", None, items), ("BOM/BOM_원본.xlsx", "No", None, bom), ("라우팅/라우팅_원본.xlsx", "순서", None, route),
                                      ("작업장/작업장_원본.xlsx", "No", ["작업장목록"], ws), ("작업장/작업장_원본.xlsx", "작업장명", ["근무시간"], C["작업장_근무시간_정답"]),
                                      ("비가동사유/비가동사유_원본.xlsx", "코드", None, codes), ("MPS/MPS_원본.xlsx", "No", ["MPS목록"], mps),
                                      ("생산계획/생산계획_원본.xlsx", "계획번호", ["생산계획"], plans), ("생산계획/생산계획_원본.xlsx", "계획번호", ["반제품소요"], subs),
                                      ("작업지시서/작업지시서_원본.xlsx", "작업지시번호", ["작업지시"], wos), ("작업지시서/작업지시서_원본.xlsx", "작업지시번호", ["소요자재"], wit),
                                      ("작업지시서/작업지시서_원본.xlsx", "작업지시번호", ["공정"], wop), ("작업실적/작업실적_원본.xlsx", "NO", None, log),
                                      ("비가동기록/비가동기록_원본.xlsx", "No", None, dt), ("4M변경점/4M변경점_원본.xlsx", "라인명", None, m4)]:
        nx, gap = _xlsx_rows(out / path, label, sheets)
        if gap:
            E.append(f"표 머리글 바로 아래가 빈 줄: {path} 시트 {gap[:3]}")
        if nx != len(rows):
            E.append(f"원본↔정답 줄 수 다름: {path} {sheets or ''} {nx} != {len(rows)}")
    # (2) 품목·BOM
    fg = [c for c, r in items.items() if r["품목구분"] == "완제품"]
    for b in bom:
        for k in ("parent", "child"):
            if b[k] not in items:
                E.append(f"BOM {b['product']} {k} {b[k]} 품목마스터에 없음")
        if b["child"] in items and items[b["child"]]["사용여부"] != "Y":
            E.append(f"BOM 이 사용중지 품목 {b['child']} 을 씀")
        if float(b["qty"]) <= 0:
            E.append(f"BOM 소요량 0 이하 {b['child']}")
    for f_ in fg:
        if not any(b["product"] == f_ for b in bom):
            E.append(f"완제품 {f_} BOM 없음")
        if not any(r["product"] == f_ for r in route):
            E.append(f"완제품 {f_} 라우팅 없음")
        cur = None
        for b in [x for x in bom if x["product"] == f_]:     # 정전개: 2단계 줄의 모품목은 바로 앞 1단계 품목
            if b["level"] == "1":
                cur = b["child"]
            elif b["parent"] != cur:
                E.append(f"BOM {f_} 정전개 순서 어긋남 no {b['no']}")
    # (3) 라우팅·작업장
    for r in route:
        if r["workcenter"] not in ws or ws[r["workcenter"]]["disabled"] == "Y":
            E.append(f"라우팅 작업장 {r['workcenter']} 없음/사용중지")
    for b in bom:
        if not any(r["product"] == b["product"] and r["name"] == b["process"] for r in route):
            E.append(f"BOM 투입공정 {b['process']} 가 라우팅에 없음 ({b['product']} {b['child']})")
    for r in ws.values():
        if r["parent"] and r["parent"] not in ws:
            E.append(f"작업장 상위 {r['parent']} 없음")
    # (4) MPS ↔ 생산계획 ↔ 작업지시 ↔ 소요자재
    if len(mps) != len(plans):
        E.append(f"MPS {len(mps)}줄 != 생산계획 {len(plans)}줄")
    for p in plans:
        mm = [m for m in mps if m["item"] == p["item"] and m["delivery"] == p["end"]]
        if len(mm) != 1 or int(mm[0]["qty"]) != int(p["qty"]):
            E.append(f"생산계획 {p['plan_no']} {p['item']} 이 MPS 와 안 맞음")
        got = sum(int(w["qty"]) for w in wos.values() if w["plan_no"] == p["plan_no"] and w["item"] == p["item"])
        if got != int(p["qty"]):
            E.append(f"{p['plan_no']} {p['item']} 지시수량 합 {got} != 계획 {p['qty']}")
        pk = sum(int(o["done"]) for o in wop if o["code"] == "OP-PKG" and wos[o["wo"]]["plan_no"] == p["plan_no"] and wos[o["wo"]]["item"] == p["item"])
        if int(p["produced"]) != pk or int(p["remain"]) != int(p["qty"]) - pk:
            E.append(f"{p['plan_no']} {p['item']} 생산수량·잔량이 포장 실적과 다름")
    for m in mps:
        if m["item"] not in items or date.fromisoformat(m["release"]) > date.fromisoformat(m["delivery"]):
            E.append(f"MPS {m['item']} {m['delivery']} 키/날짜 이상")
    for s in subs:
        bl = [b for b in bom if b["product"] == s["parent"] and b["child"] == s["sub"] and b["level"] == "1"]
        pl = [p for p in plans if p["plan_no"] == s["plan_no"] and p["item"] == s["parent"]]
        if not bl or not pl:
            E.append(f"반제품소요 {s['plan_no']} {s['sub']} 키 끊김")
        elif int(s["need"]) != int(float(bl[0]["qty"])) * int(pl[0]["qty"]):
            E.append(f"반제품소요 {s['sub']} 소요량 {s['need']} != BOM×계획")
    for w in wos.values():
        if w["item"] not in items or not any(p["plan_no"] == w["plan_no"] and p["item"] == w["item"] for p in plans):
            E.append(f"{w['wo']} 키 끊김")
        if int(w["qty"]) > LOT_MAX:
            E.append(f"{w['wo']} LOT {w['qty']} > {LOT_MAX}")
        if w["issue"] > w["start"]:
            E.append(f"{w['wo']} 지시일이 시작보다 늦음")
    for r in wit:
        b = [x for x in bom if x["product"] == wos[r["wo"]]["item"] and x["child"] == r["code"] and x["level"] == "1"]
        if not b:
            E.append(f"소요자재 {r['wo']} {r['code']} 가 BOM 1단계에 없음")
        elif int(r["need"]) != int(float(b[0]["qty"])) * int(wos[r["wo"]]["qty"]):
            E.append(f"소요자재 {r['wo']} {r['code']} 필요수량 틀림")
        if not (int(r["used"]) <= int(r["issued"]) <= int(r["need"])):
            E.append(f"소요자재 {r['wo']} {r['code']} 사용≤불출≤필요 위반")
    # (5) 공정 일정·상태
    by_wo = defaultdict(list)
    for o in wop:
        by_wo[o["wo"]].append(o)
        s_, e_ = datetime.fromisoformat(o["start"]), datetime.fromisoformat(o["end"])
        if o["wc"] not in ws:
            E.append(f"공정 작업장 {o['wc']} 없음")
        if e_ <= s_ or next_work(s_) != s_:
            E.append(f"{o['wo']} {o['name']} 시각 이상 {o['start']}~{o['end']}")
        want = "완료" if e_ <= T else "진행" if s_ < T else "대기"
        if o["status"] != want:
            E.append(f"{o['wo']} {o['name']} 상태 {o['status']} != {want}")
        if o["status"] == "대기" and (int(o["done"]) or int(o["loss"])):
            E.append(f"{o['wo']} {o['name']} 대기인데 실적 있음")
        if int(o["done"]) + int(o["loss"]) > int(o["qty_in"]):
            E.append(f"{o['wo']} {o['name']} 처리수량 > 투입")
    for wo, lst in by_wo.items():
        lst.sort(key=lambda x: int(x["seq"]))
        rt = sorted([r for r in route if r["product"] == wos[wo]["item"]], key=lambda r: int(r["seq"]))
        if [x["name"] for x in lst] != [r["name"] for r in rt]:
            E.append(f"{wo} 공정 목록이 라우팅과 다름")
        for a, b in zip(lst, lst[1:]):
            if b["start"] < a["end"]:
                E.append(f"{wo} 공정 순서 위반 {a['name']}→{b['name']}")
            if a["status"] == "완료" and int(b["qty_in"]) != int(a["done"]):
                E.append(f"{wo} {b['name']} 투입 {b['qty_in']} != 앞 공정 양품 {a['done']}")
        st = "완료" if all(x["status"] == "완료" for x in lst) else "미시작" if all(x["status"] == "대기" for x in lst) else "진행중"
        if wos[wo]["status"] != st:
            E.append(f"{wo} 지시 상태 {wos[wo]['status']} != {st}")
        if st == "완료" and int(wos[wo]["produced"] or -1) != int(lst[-1]["done"]):
            E.append(f"{wo} 생산수량 != 포장 양품")
    ev = defaultdict(list)
    for o in wop:
        ev[o["wc"]] += [(datetime.fromisoformat(o["start"]), 1), (datetime.fromisoformat(o["end"]), -1)]
    for wc, pts in ev.items():
        cur = 0
        for t, d in sorted(pts, key=lambda x: (x[0], x[1])):
            cur += d
            if cur > int(ws[wc]["parallel"]):
                E.append(f"{wc} 동시 작업 {cur} > 동시처리 {ws[wc]['parallel']} @ {t}")
                break
    # (6) 작업실적
    agg = defaultdict(lambda: [0, 0])
    for r in log:
        if datetime.fromisoformat(r["date"] + " " + r["end"]) > T:
            E.append(f"작업실적 {r['date']} {r['no']} 이 출력 기준 시각 뒤")
        if r["wo"] not in wos or r["workcenter"] not in ws or r["item"] != wos.get(r["wo"], {}).get("item"):
            E.append(f"작업실적 {r['date']} {r['no']} 키 끊김")
        if int(r["good_qty"]) + int(r["loss"]) != int(r["qty"]) or int(r["qty"]) <= 0:
            E.append(f"작업실적 {r['date']} {r['no']} 수량 이상")
        agg[(r["wo"], r["process"])][0] += int(r["good_qty"])
        agg[(r["wo"], r["process"])][1] += int(r["loss"])
    for o in wop:
        if o["status"] != "대기" and tuple(agg[(o["wo"], o["name"])]) != (int(o["done"]), int(o["loss"])):
            E.append(f"{o['wo']} {o['name']} 공정 실적 {o['done']}/{o['loss']} != 일지 합 {agg[(o['wo'], o['name'])]}")
    spec = {s["model"]: s for s in json.loads((out / "사양추출.json").read_text(encoding="utf-8"))}
    sn = defaultdict(list)
    for r in log:
        if r["serial_from"]:
            a, b = int(r["serial_from"][-5:]), int(r["serial_to"][-5:])
            if b - a + 1 != int(r["good_qty"]):
                E.append(f"시리얼 개수 != 양품 {r['date']} {r['no']}")
            sn[r["item"]].append((a, b))
        s = spec.get(r["item"])
        if r["power_w"] and s and s["power"] and float(r["power_w"]) > s["power"][0]["v"] and "기준" not in r["note"]:
            E.append(f"작업실적 {r['date']} {r['no']} 소비전력 기준 초과인데 비고 없음")
        if r["weight_g"] and s and s.get("weight_g"):
            wt = s["weight_g"]["v"]
            if abs(float(r["weight_g"]) - wt) > wt * 0.03 and "이탈" not in r["note"]:
                E.append(f"작업실적 {r['date']} {r['no']} 중량 이탈인데 비고 없음")
    for it, lst in sn.items():
        lst.sort()
        for (a1, b1), (a2, b2) in zip(lst, lst[1:]):
            if a2 <= b1:
                E.append(f"{it} 시리얼 겹침 {a1}-{b1} / {a2}-{b2}")
    # (7) 비가동
    busy = defaultdict(list)
    for r in dt:
        s_, e_ = datetime.fromisoformat(r["start"]), datetime.fromisoformat(r["end"])
        if r["code"] not in codes or codes[r["code"]]["use"] != "Y":
            E.append(f"비가동 {r['start']} 사유코드 {r['code']} 없음/사용중지")
        if e_ > T or int(r["minutes"]) != int((e_ - s_).total_seconds() // 60) or segments(s_, e_) != [(s_, e_)]:
            E.append(f"비가동 {r['start']} 시간 이상")
        if not any(o["wo"] == r["wo"] and o["wc"] == r["wc"] and o["start"] <= r["start"] and r["end"] <= o["end"] for o in wop):
            E.append(f"비가동 {r['wc']} {r['start']} 이 해당 작업 시간 밖")
        if any(not (e_ <= a or s_ >= b) for a, b in busy[r["wc"]]):
            E.append(f"비가동 {r['wc']} {r['start']} 겹침")
        busy[r["wc"]].append((s_, e_))
    # (8) 4M
    lines_used = {r["workcenter"] for r in route}
    for d_ in workdays(FIRST_MONDAY, T.date()):
        got = {r["line"] for r in m4 if r["date"] == ds(d_)}
        if got != lines_used:
            E.append(f"4M {ds(d_)} 라인 누락 {lines_used - got}")
    for r in m4:
        if r["date"] > ds(T.date()):
            E.append(f"4M {r['date']} 이 출력 기준일 뒤 (미래 기록)")
        if r["change"] != "없음" and not r["m4"]:
            E.append(f"4M {r['date']} {r['line']} 4M 구분 없음")
    # (9) 근거 칸이 비지 않았나
    for n_, rows in C.items():
        if rows and "basis" in rows[0]:
            miss = sum(1 for r in rows if not r["basis"].strip())
            if miss:
                E.append(f"{n_}: basis 빈 줄 {miss}")
    # (10) 틀 고정(시트 끝까지 이어지는 선)이 없는지
    import zipfile
    for f in FOLDERS:
        with zipfile.ZipFile(out / f / f"{f}_원본.xlsx") as z:
            if any(b"<pane " in z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/sheet")):
                E.append(f"{f} 원본에 틀 고정이 있음 (표 밖으로 선이 삐져나옴)")
    # (11) 원본(현업 모양)에 근거 꼬리표가 새지 않았나
    for f in FOLDERS:
        wb = openpyxl.load_workbook(out / f / f"{f}_원본.xlsx", read_only=True)
        if any(isinstance(v, str) and ("[사양서" in v or "[가상]" in v) for wsx in wb.worksheets for row in wsx.iter_rows(values_only=True) for v in row):
            E.append(f"{f} 원본에 근거 꼬리표가 들어감")
        wb.close()
    # (12) PDF·docx·hwpx 표 칸 글자가 xlsx 와 같은지
    for f in EXPORT_DOCS if formats else []:
        E += [f"{f} {e}" for e in check_formats(out / f / f"{f}_원본.xlsx", formats)]
    st = defaultdict(int)
    for w in wos.values():
        st[w["status"]] += 1
    notes = [f"품목 {len(items)} / BOM {len(bom)} / 공정 {len(route)} / 작업장 {len(ws)} / 사유 {len(codes)} / MPS {len(mps)} / 계획 {len(plans)} / 지시 {len(wos)} / "
             f"일지 {len(log)} / 비가동 {len(dt)} / 4M {len(m4)}", "지시 상태 " + ", ".join(f"{k} {v}" for k, v in st.items())]
    return E, notes


def _check_pdf(p):
    """넣은 파일이 읽을 수 있는 PDF 인지 미리 본다. 아니면 오류 화면 대신 알아들을 수 있는 말로 멈춘다."""
    p = Path(p)
    if not p.exists():
        raise SystemExit(f"파일을 찾지 못했습니다: {p}\n  경로가 맞는지, 파일 이름에 오타가 없는지 확인하세요.")
    if p.is_dir():
        raise SystemExit(f"폴더가 아니라 PDF 파일을 넣어 주세요: {p}")
    if p.suffix.lower() != ".pdf":
        raise SystemExit(f"PDF 파일만 넣을 수 있습니다: {p.name}")
    try:
        with pdfplumber.open(p) as d:
            text = "".join((pg.extract_text() or "") for pg in d.pages[:2])
    except Exception:
        raise SystemExit(f"PDF 로 열 수 없습니다 (깨졌거나 PDF 가 아닌 파일): {p.name}")
    if len(text.strip()) < 50:
        raise SystemExit(f"글자를 읽을 수 없는 PDF 입니다 (스캔·이미지 문서로 보임): {p.name}\n  글자가 들어 있는 제품 사양서 PDF 를 넣어 주세요.")


def main():
    ap = argparse.ArgumentParser(description="제품 사양서 PDF 로 생산관리 가상 문서 11종을 만든다.")
    ap.add_argument("pdfs", nargs="+", help="제품 사양서(데이터시트) PDF. 여러 개면 한 공장에서 같이 만드는 것으로 본다")
    ap.add_argument("--out", help="저장 폴더 (없으면 지금 위치에 '생산관리문서_<모델>' 폴더를 만든다)")
    ap.add_argument("--seed", type=int, default=0, help="가상 값을 다르게 뽑고 싶을 때 바꾸는 번호 (같은 번호면 같은 결과)")
    ap.add_argument("--formats", default=",".join(OUT_FORMATS), help="xlsx 말고 더 낼 형식 (쉼표로, 빈 값이면 xlsx 만): pdf,docx,hwpx")
    ap.add_argument("--category", choices=sorted(CATEGORY_NAME), help="제품군을 직접 정한다. 없으면 폴더 이름, 그다음 사양서 내용으로 정한다")
    a = ap.parse_args()
    formats = tuple(x.strip().lower() for x in a.formats.split(",") if x.strip())
    bad = [x for x in formats if x not in OUT_FORMATS]
    if bad:
        raise SystemExit(f"모르는 형식입니다: {', '.join(bad)}\n  쓸 수 있는 형식: {', '.join(OUT_FORMATS)} (xlsx 는 언제나 만든다)")
    need = {"docx": "docx", "pdf": "pymupdf"}                # docx 는 만들 때, pymupdf 는 PDF 를 검사할 때 쓴다
    missing = sorted({mod for f, mod in need.items() if f in formats and importlib.util.find_spec(mod) is None})
    if missing:
        names = {"docx": "python-docx", "pymupdf": "pymupdf"}
        raise SystemExit(f"필요한 파이썬 패키지가 없습니다: {', '.join(names[m] for m in missing)}\n  설치: pip install {' '.join(names[m] for m in missing)}"
                         "\n  (또는 --formats 에서 그 형식을 빼고 실행)")
    if "pdf" in formats and not find_soffice():
        print("  ! LibreOffice 를 찾지 못해 PDF 는 건너뜁니다. PDF 도 필요하면 LibreOffice 를 설치하거나 SOFFICE 환경변수에 soffice 경로를 적으세요.")
        formats = tuple(x for x in formats if x != "pdf")
    specs = []
    for p in a.pdfs:
        _check_pdf(p)
        found = read_specs(p, a.category)
        if not found:
            raise SystemExit(f"제품 사양서로 보기 어렵습니다: {Path(p).name} (모델명을 찾지 못함). 한화비전 제품 사양서(데이터시트) PDF 를 넣어 주세요.")
        for s in found:
            if not (s["sensor"] or s["thermal"]) or not (s["resolution"] or s["power"]):
                raise SystemExit(f"제품 사양서로 보기 어렵습니다: {Path(p).name} (촬상 소자·해상도·소비전력을 찾지 못함). 시험성적서·안내문은 넣지 마세요.")
            specs.append(s)
    name = specs[0]["model"] if len(specs) == 1 else f"{specs[0]['model']}_외{len(specs) - 1}"
    out = Path(a.out) if a.out else Path.cwd() / f"생산관리문서_{name}"
    build(specs, out, a.seed, formats)
    errs, notes = verify(out, formats)
    print(f"== {', '.join(s['model'] for s in specs)} -> {out}")
    for n_ in notes:
        print("  " + n_)
    print("  검증 오류 없음" if not errs else f"  검증 오류 {len(errs)}건")
    for e in errs[:25]:
        print("   -", e)
    raise SystemExit(1 if errs else 0)


if __name__ == "__main__":
    main()
